"""Militaires Sans Frontières — primera capa de asignación de rutas."""
from __future__ import annotations
import csv, io, os, re, unicodedata, urllib.request
from datetime import date, datetime, time as datetime_time
from pathlib import Path
import openpyxl
import streamlit as st
from mother_base_theme import render_system_stamp

def _inject_style():
    st.markdown('<style>.msf-shell{border:3px solid #111;background:#fffdf5;padding:28px;box-shadow:8px 8px 0 #111;margin-bottom:28px}.msf-label{display:inline-block;background:#5f7cff;color:#111;border:3px solid #111;padding:8px 12px;font-weight:800;letter-spacing:.08em}.msf-note{border:2px solid #111;background:#e8edff;padding:14px;margin-top:16px}</style>', unsafe_allow_html=True)

def _secret(name):
    def find(value):
        if isinstance(value, dict):
            for key,item in value.items():
                if str(key).strip().upper()==name.upper() and item not in (None,''):
                    return str(item).strip()
                found=find(item)
                if found: return found
        return ''
    try:
        found=find(st.secrets.to_dict() if hasattr(st.secrets,'to_dict') else st.secrets)
        if found: return found
    except Exception:
        pass
    return os.environ.get(name,'').strip()


def _cell(row,index,default=None):
    """Lee filas de hojas externas que pueden venir truncadas al final."""
    return row[index] if index < len(row) else default

def _read_csv_dispatchtrack(uploaded_file):
    rows=list(csv.reader(io.StringIO(uploaded_file.getvalue().decode('utf-8-sig'))))
    if len(rows)<2: raise ValueError('El CSV debe tener marcadores y encabezados.')
    markers,headers=rows[0],rows[1]
    selected=[i for i,v in enumerate(markers) if v.strip()=='DispatchTrack']
    if not selected: raise ValueError('No encontré columnas marcadas como DispatchTrack en la fila 1.')
    names=[headers[i].strip() or f'COLUMNA_{i+1}' for i in selected]
    return names,[{n:(row[i].strip() if i<len(row) else '') for n,i in zip(names,selected)} for row in rows[2:] if any(x.strip() for x in row)]

def _read_excel_dispatchtrack(uploaded_file):
    raw=uploaded_file.getvalue(); wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    try: rows=[list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    finally: wb.close(); del raw
    if not rows: raise ValueError('El Excel está vacío.')
    markers=[str(v or '').strip() for v in rows[0]]
    if 'DispatchTrack' in markers and len(rows)>=2: header,selected,data=rows[1],[i for i,v in enumerate(markers) if v=='DispatchTrack'],rows[2:]
    else: header,selected,data=rows[0],list(range(len(rows[0]))),rows[1:]
    names=[str(header[i] or f'COLUMNA_{i+1}').strip() for i in selected]
    return names,[{n:str(row[i] if i<len(row) and row[i] is not None else '').strip() for n,i in zip(names,selected)} for row in data if any(row)]

def _read_dispatchtrack_file(uploaded_file):
    return _read_csv_dispatchtrack(uploaded_file) if Path(uploaded_file.name).suffix.lower()=='.csv' else _read_excel_dispatchtrack(uploaded_file)

def _load_logistics_book():
    sheet_id=_secret('DATA_LOGISTICS_ID')
    if not sheet_id: raise ValueError('Falta DATA_LOGISTICS_ID en Secrets.')
    with urllib.request.urlopen(f'https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx',timeout=60) as response: raw=response.read()
    return raw

def _load_fleet():
    raw=_load_logistics_book(); wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    try:
        if 'FLOTA' not in wb.sheetnames: raise ValueError('DATA_LOGISTICS no contiene la hoja FLOTA.')
        rows=list(wb['FLOTA'].iter_rows(values_only=True))
    finally: wb.close(); del raw
    headers=[str(v or '').strip().upper() for v in rows[0]]; required={'UNIDAD','CANTIDAD','CAPACIDAD UNO','VUELTAS'}
    if not required.issubset(headers): raise ValueError('FLOTA debe contener UNIDAD, CANTIDAD, CAPACIDAD UNO y VUELTAS.')
    pos={h:headers.index(h) for h in required}; fleet=[]
    for row in rows[1:]:
        try: qty=int(float(_cell(row,pos['CANTIDAD'],0) or 0)); turns=int(float(_cell(row,pos['VUELTAS'],1) or 1)); cap=float(_cell(row,pos['CAPACIDAD UNO'],0) or 0)
        except (TypeError,ValueError): continue
        if qty>0 and turns>0 and cap>0:
            fleet.append({'UNIDAD':str(_cell(row,pos['UNIDAD'],'')).strip(),'CANTIDAD':qty,'CAPACIDAD':cap,'VUELTAS':turns})
    if not fleet: raise ValueError('FLOTA no tiene unidades válidas.')
    return sorted(fleet,key=lambda x:float(x['CAPACIDAD']),reverse=True)

def _load_directions():
    raw=_load_logistics_book(); wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    try:
        if 'DIRECCIONES' not in wb.sheetnames: return {}
        rows=list(wb['DIRECCIONES'].iter_rows(values_only=True))
    finally: wb.close(); del raw
    if not rows: return {}
    headers=[str(v or '').strip().upper() for v in rows[0]]; pos={h:headers.index(h) for h in headers}
    if 'WAREHOUSE_ID' not in pos: return {}
    result={}
    for row in rows[1:]:
        key=str(_cell(row,pos['WAREHOUSE_ID'],'') or '').strip().replace('.0','')
        data={'CITY':str(_cell(row,pos['CITY'],'')).strip() if 'CITY' in pos and _cell(row,pos['CITY']) is not None else '', 'CONSTRAIN':str(_cell(row,pos['CONSTRAIN'],'')).strip() if 'CONSTRAIN' in pos and _cell(row,pos['CONSTRAIN']) is not None else '', 'PRIORITY':_number(_cell(row,pos['PRIORITY'])) if 'PRIORITY' in pos else 0}
        if key: result[key]=data
        if 'WAREHOUSE_NAME' in pos and _cell(row,pos['WAREHOUSE_NAME']): result[_norm(_cell(row,pos['WAREHOUSE_NAME']))]=data
    return result

def _norm(value):
    return ''.join(ch for ch in unicodedata.normalize('NFKD',str(value or '')) if not unicodedata.combining(ch)).strip().casefold()

def _number(value):
    if value is None: return 0.0
    text=str(value).strip().replace(',','.')
    if not text or text.lower() in {'none','nan','null'}: return 0.0
    try: return float(text)
    except (TypeError,ValueError): return 0.0


def _round_pallets(value):
    value=_number(value)
    if value<=0: return 0
    whole=int(value); return max(1,whole+(1 if value-whole>=0.3 else 0))


def _split_destination_address(address, split_index):
    """Crea una dirección distinguible para una parte adicional del mismo pedido."""
    value=str(address or '')
    if split_index<=0 or not value:
        return value
    match=re.search(r'(,\s*)(\d+)([A-Za-z]*)',value)
    if not match:
        return value
    number=str(int(match.group(2))+split_index)
    return f'{value[:match.start()]}{match.group(1)}{number}{match.group(3)}{value[match.end():]}'


def _mark_split_destination(row,headers,split_index):
    """Marca las partes derivadas para que DispatchTrack las trate como destinos únicos."""
    if split_index<=0:
        return
    normalized=lambda header:_norm(header).replace(' ','')
    contact_headers=[
        header for header in headers
        if normalized(header).replace('*','') in {'identificadorcontacto','nombrecontacto'}
    ]
    address_header=next(
        (header for header in headers if normalized(header).replace('*','')=='direccion'),None
    )
    for header in contact_headers:
        value=str(row.get(header,'') or '').strip()
        if value:
            row[header]=f'{value}-{split_index}'
    if address_header:
        row[address_header]=_split_destination_address(row.get(address_header,''),split_index)


def _display_number(value):
    """Serializa una cantidad sin introducir ruido de punto flotante."""
    return f'{float(value):.12f}'.rstrip('0').rstrip('.') or '0'


def _apply_split_quantities(clone, row, cap_header, cap2_header, cost_header, part_capacity, original_capacity):
    """Conserva los totales de capacidad, capacidad secundaria y costo al partir una fila."""
    clone[cap_header]=_display_number(part_capacity)
    if original_capacity<=0:
        return
    ratio=part_capacity/original_capacity
    for header in (cap2_header,cost_header):
        if not header or str(row.get(header,'')).strip()=='':
            continue
        clone[header]=_display_number(_number(row.get(header,''))*ratio)


def _priority_store_key(info):
    """Agrupa una tienda física aunque DispatchTrack haya agregado -1, -2, etc."""
    return re.sub(r'-\d+$','',_norm(info.get('_STORE','')))


def _select_first_turn_torton_stores(ordered, cap_header, first_turn_slots):
    """Maximiza tiendas PRIORITY=1 completas; después, uso de slots Torton."""
    stores={}
    for row,info in ordered:
        if _number(info.get('PRIORITY',0))!=1 or _norm(info.get('CONSTRAIN',''))=='1.5':
            continue
        key=_priority_store_key(info)
        slots=max(_round_pallets(row.get(cap_header,'')),5 if _number(row.get(cap_header,''))>4 else 0)
        stores[key]=stores.get(key,0)+slots
    # Programación dinámica: primero más tiendas completas, después más uso de Torton.
    choices={0:()}
    for key,slots in sorted(stores.items()):
        if slots>first_turn_slots:
            continue
        for used,selected in list(choices.items())[::-1]:
            total=used+slots
            if total>first_turn_slots:
                continue
            candidate=selected+(key,)
            current=choices.get(total)
            if current is None or len(candidate)>len(current) or (len(candidate)==len(current) and candidate<current):
                choices[total]=candidate
    best=max(choices.items(),key=lambda item:(len(item[1]),item[0],tuple(reversed(item[1]))))
    return set(best[1])


def _assign_priority_torton_labels(ordered,headers,cap_header,selected_stores):
    """Asigna tiendas P1 completas a Torton físico, agrupando por cercanía."""
    lat_header=next((h for h in headers if _norm(h) in {'latitud','latitude'}),None)
    lon_header=next((h for h in headers if _norm(h) in {'longitud','longitude'}),None)
    groups={}
    for row,info in ordered:
        key=_priority_store_key(info)
        if key not in selected_stores:
            continue
        group=groups.setdefault(key,{'slots':0,'points':[]})
        group['slots']+=max(_round_pallets(row.get(cap_header,'')),5 if _number(row.get(cap_header,''))>4 else 0)
        if lat_header and lon_header:
            lat,lon=_number(row.get(lat_header,'')),_number(row.get(lon_header,''))
            if lat and lon: group['points'].append((lat,lon))
    labels={f'TORTON_{index}':{'free':8,'points':[]} for index in range(1,4)}
    assigned={}
    # Primero las tiendas grandes; entre opciones factibles elige la zona más cercana.
    for key,group in sorted(groups.items(),key=lambda item:(-item[1]['slots'],item[0])):
        eligible=[(label,state) for label,state in labels.items() if state['free']>=group['slots']]
        if not eligible:
            continue
        point=(sum(x for x,_ in group['points'])/len(group['points']),sum(y for _,y in group['points'])/len(group['points'])) if group['points'] else None
        def score(item):
            label,state=item
            if point and state['points']:
                center=(sum(x for x,_ in state['points'])/len(state['points']),sum(y for _,y in state['points'])/len(state['points']))
                return ((point[0]-center[0])**2+(point[1]-center[1])**2,label)
            return (-state['free'],label)
        label,state=min(eligible,key=score)
        state['free']-=group['slots']; state['points'].extend(group['points']); assigned[key]=label
    return assigned


def _assign_and_split(headers,rows,fleet,directions=None,allowed_cities=None,*,respect_constraints=True,split_orders=True,split_capacity=4):
    directions=directions or {}; cap_header=next((h for h in headers if h.upper()=='CAPACIDAD UNO'),None)
    id_header=next((h for h in headers if h.upper() in {'CÓDIGO DE PARADA*','CODIGO DE PARADA*','TRANSFER ID'}),headers[0])
    origin_header=next((h for h in headers if h.upper() in {'IDENTIFICADOR CONTACTO*','IDENTIFICADOR CONTACTO','WAREHOUSE_NAME'}),None) or next((h for h in headers if h.upper() in {'CT ORIGEN*','CT ORIGEN','WAREHOUSE_ID'}),None)
    type_header=next((h for h in headers if _norm(h).replace(' ','') in {'tipocamion','tipodecamion'}),None)
    if not cap_header: raise ValueError('El archivo no tiene CAPACIDAD UNO.')
    bins=[]
    for item in fleet:
        for turn in range(1,int(item.get('VUELTAS',1))+1):
            for _ in range(int(item['CANTIDAD'])):
                bins.append({'UNIDAD':item['UNIDAD'],'CAPACIDAD':float(item['CAPACIDAD']),'LIBRE':float(item['CAPACIDAD']),'ORIGEN':None,'VUELTA':turn})
    importance_header=next((h for h in headers if _norm(h).replace(' ','')=='importancia'),None)
    cap2_header=next((h for h in headers if _norm(h).replace(' ','')=='capacidaddos'),None)
    cost_header=next((h for h in headers if _norm(h).replace(' ','')=='costoitem'),None)
    output=[]; warnings=[]; split_counts={}
    ordered=[]
    for row in rows:
        origin=str(row.get(origin_header,'')).strip() if origin_header else ''
        info=dict(directions.get(origin.replace('.0',''),directions.get(_norm(origin),{})))
        info['_STORE']=_norm(origin)
        if allowed_cities and _norm(info.get('CITY','')) not in {_norm(city) for city in allowed_cities}:
            continue
        ordered.append((row,info))
    if not ordered:
        ordered=[(row,{}) for row in rows]
    ordered.sort(key=lambda x:(0 if _number(x[1].get('PRIORITY',0))==1 else 1, _priority_store_key(x[1]), -_round_pallets(x[0].get(cap_header,''))))
    first_turn_slots=sum(int(item['CANTIDAD'])*int(item.get('VUELTAS',1)>=1)*int(item['CAPACIDAD']) for item in fleet if _norm(item['UNIDAD'])=='torton')
    first_turn_priority_stores=_select_first_turn_torton_stores(ordered,cap_header,first_turn_slots)
    priority_torton_labels=_assign_priority_torton_labels(ordered,headers,cap_header,first_turn_priority_stores)
    for row,info in ordered:
        original_capacity=_number(row.get(cap_header,'')); pallets=_round_pallets(original_capacity); pallets=max(pallets, 5 if original_capacity>4 else pallets); left=pallets; original_left=original_capacity
        allowed=_norm(info.get('CONSTRAIN','')) if respect_constraints else ''
        split_limit=max(1,int(split_capacity))
        must_split=respect_constraints and allowed=='1.5'
        can_split=split_orders or must_split
        while left>0:
            priority_one=_number(info.get('PRIORITY',0))==1
            priority_torton_label=priority_torton_labels.get(_priority_store_key(info))
            priority_torton=priority_one and allowed!='1.5' and bool(priority_torton_label)
            if priority_one and allowed=='1.5':
                candidates=[b for b in bins if b['LIBRE']>=1 and _norm(b['UNIDAD'])=='1.5']
                candidates.sort(key=lambda b:(0 if b.get('ORIGEN')==info.get('_STORE') else 1,b['LIBRE']))
            elif priority_torton:
                candidates=[b for b in bins if b['LIBRE']>=1 and _norm(b['UNIDAD'])=='torton' and b['VUELTA']==1]
                candidates.sort(key=lambda b:(0 if b.get('ORIGEN')==info.get('_STORE') else 1,b['LIBRE']))
            elif priority_one:
                candidates=[b for b in bins if b['LIBRE']>=1 and _norm(b['UNIDAD'])=='3.5']
                candidates.sort(key=lambda b:(0 if b.get('ORIGEN')==info.get('_STORE') else 1,-b['CAPACIDAD']))
            else:
                # Esta iteración reserva toda etiqueta TORTON para la primera
                # vuelta P1. El resto queda disponible para flota regular.
                candidates=[b for b in bins if b['LIBRE']>=1 and _norm(b['UNIDAD'])!='torton' and (not allowed or _norm(b['UNIDAD'])==allowed)]
                if allowed=='1.5': candidates.sort(key=lambda b:(0 if _norm(b['UNIDAD'])=='1.5' else 1,b['LIBRE']))
                else: candidates.sort(key=lambda b:(0 if b.get('ORIGEN')==info.get('_STORE') else 1,-b['CAPACIDAD']))
            if not can_split:
                candidates=[b for b in candidates if int(b['LIBRE'])>=left]
            if not candidates:
                if not can_split:
                    remainder=dict(row)
                    fallback_unit='1.5' if allowed=='1.5' else ('3.5' if priority_one else '')
                    if type_header: remainder[type_header]=fallback_unit
                    if importance_header: remainder[importance_header]='1' if fallback_unit in {'torton','1.5'} else ''
                    output.append(remainder)
                    break
                while left>0:
                    amount=min(left,2 if allowed=='1.5' else split_limit)
                    part=original_left*amount/left if left else 0
                    remainder=dict(row)
                    _apply_split_quantities(remainder,row,cap_header,cap2_header,cost_header,part,original_capacity)
                    fallback_unit='1.5' if allowed=='1.5' else ('3.5' if priority_one else '')
                    if type_header: remainder[type_header]=fallback_unit
                    if importance_header: remainder[importance_header]='1' if fallback_unit in {'torton','1.5'} else ''
                    base_id=str(row.get(id_header,'')); split_index=split_counts.get(base_id,0)
                    if split_index>0:
                        remainder[id_header]=f'{base_id}-{split_index}'
                        _mark_split_destination(remainder,headers,split_index)
                    split_counts[base_id]=split_index+1
                    output.append(remainder); original_left=max(0,original_left-part); left-=amount
                break
            bin_=candidates[0]; chunk_limit=2 if _norm(bin_['UNIDAD'])=='1.5' and respect_constraints else (split_limit if can_split else left); amount=min(left,chunk_limit,int(bin_['LIBRE']))
            clone=dict(row)
            base_id=str(row.get(id_header,'')); split_index=split_counts.get(base_id,0)
            if split_index>0:
                clone[id_header]=f"{base_id}-{split_index}"
                _mark_split_destination(clone,headers,split_index)
            split_counts[base_id]=split_index+1
            part=original_left*amount/left if left else 0
            _apply_split_quantities(clone,row,cap_header,cap2_header,cost_header,part,original_capacity)
            original_left=max(0,original_left-part)
            assigned_unit=_norm(bin_['UNIDAD'])
            if type_header:
                clone[type_header]=priority_torton_label if priority_torton else (bin_['UNIDAD'] if assigned_unit in {'torton','1.5'} or (priority_one and assigned_unit=='3.5') else '')
            visible_type=clone.get(type_header,'') if type_header else ''
            if importance_header: clone[importance_header]='1' if _norm(visible_type) in {'torton','1.5'} else ''
            output.append(clone); bin_['ORIGEN']=info.get('_STORE'); bin_['LIBRE']-=amount; left-=amount
    return output,warnings

DISPATCHTRACK_OUTPUT_HEADERS=(
    'CÓDIGO DE PARADA*','LATITUD','LONGITUD','DIRECCION*','NOMBRE ITEM*',
    'CANTIDAD*','CODIGO ITEM','FECHA MIN ENTREGA*','FECHA MAX ENTREGA*',
    'MIN VENTANA HORARIA 1*','MAX VENTANA HORARIA 1*','MIN VENTANA HORARIA 2',
    'MAX VENTANA HORARIA 2','COSTO ITEM','CAPACIDAD UNO','CAPACIDAD DOS',
    'SERVICE TIME','IMPORTANCIA','IDENTIFICADOR CONTACTO*','NOMBRE CONTACTO*',
    'TELEFONO','EMAIL CONTACTO','CT ORIGEN*','Tipo camion','SERVICE TIME MINIMO',
)

EXCEL_TEXT_FIELD_NAMES=frozenset({
    'fechaminentrega',
    'fechamaxentrega',
    'minventanahoraria1',
    'minventanahorario1',
    'maxventanahoraria1',
    'maxventanahorario1',
})


def _field_name(header):
    name=_norm(header).replace(' ','').replace('*','')
    # Algunos exports antiguos escriben HORARIO en vez de HORARIA; la plantilla
    # aceptada por DispatchTrack usa HORARIA y el valor pertenece al mismo campo.
    return {'minventanahorario1':'minventanahoraria1','maxventanahorario1':'maxventanahoraria1'}.get(name,name)


def _dispatchtrack_text_value(header,value):
    """Normaliza fecha/hora y las conserva como texto en el XLSX de salida."""
    if value is None:
        return ''
    text=str(value)
    field_name=_field_name(header)
    if field_name not in EXCEL_TEXT_FIELD_NAMES:
        return value
    text=text.lstrip("'").strip()
    if not text:
        return ''
    if field_name.startswith('fecha'):
        if isinstance(value,(datetime,date)):
            return value.strftime('%d/%m/%Y')
        for pattern in ('%d/%m/%Y','%Y-%m-%d','%Y-%m-%d %H:%M:%S'):
            try:
                return datetime.strptime(text,pattern).strftime('%d/%m/%Y')
            except ValueError:
                pass
    else:
        if isinstance(value,(datetime,datetime_time)):
            return value.strftime('%H:%M')
        for pattern in ('%H:%M','%H:%M:%S','%Y-%m-%d %H:%M:%S'):
            try:
                return datetime.strptime(text,pattern).strftime('%H:%M')
            except ValueError:
                pass
    return text


def _to_xlsx(headers,rows):
    """Entrega el XLSX con el encabezado español exacto que acepta DispatchTrack."""
    source_headers={_field_name(header):header for header in headers if str(header).strip()}
    output=io.BytesIO()
    workbook=openpyxl.Workbook()
    worksheet=workbook.active
    worksheet.title='Sheet1'
    worksheet.append(DISPATCHTRACK_OUTPUT_HEADERS)
    for row in rows:
        values=[]
        for header in DISPATCHTRACK_OUTPUT_HEADERS:
            source_header=source_headers.get(_field_name(header))
            value=row.get(source_header,'') if source_header else ''
            values.append(_dispatchtrack_text_value(header,value))
        worksheet.append(values)
    for column,header in enumerate(DISPATCHTRACK_OUTPUT_HEADERS,start=1):
        if _field_name(header) in EXCEL_TEXT_FIELD_NAMES:
            for cell in worksheet.iter_cols(min_col=column,max_col=column,min_row=2):
                cell[0].number_format='@'
    worksheet.auto_filter.ref=worksheet.dimensions
    workbook.save(output)
    workbook.close()
    return output.getvalue()

def render():
    render_system_stamp('MODULE 02 / MILITAIRES SANS FRONTIÈRES'); _inject_style()
    st.markdown('<section class="mb-hero"><span class="mb-kicker">MODULE / 02 · INTELLIGENCE</span><h1>MILITAIRES SANS<br>FRONTIÈRES.</h1><p>Asignación inicial de pedidos a la flota y preparación de archivos DispatchTrack.</p></section>',unsafe_allow_html=True)
    st.markdown('<div class="msf-shell">',unsafe_allow_html=True); st.markdown('<span class="msf-label">CARGA DE PEDIDOS</span>',unsafe_allow_html=True)
    try:
        fleet_config=_load_fleet(); direction_config=_load_directions()
        cities=sorted({str(v.get('CITY','')).strip() for v in direction_config.values() if v.get('CITY')})
        units=sorted({str(v.get('UNIDAD','')).strip() for v in fleet_config if v.get('UNIDAD')})
    except Exception as exc:
        st.error(f'No se pudo leer DATA_LOGISTICS: {exc}'); fleet_config=[]; direction_config={}; cities=[]; units=[]
    with st.expander('CONFIGURACIÓN DE PLANEACIÓN',expanded=True):
        c1,c2=st.columns(2)
        selected_cities=c1.multiselect('Ciudades',cities,default=[c for c in cities if _norm(c).startswith('ciudad de')],key='msf_cities')
        selected_units=c2.multiselect('Unidades disponibles',units,default=units,key='msf_units')
        c3,c4=st.columns(2)
        use_constraints=c3.checkbox('Respetar configuraciones bloqueantes',value=True,key='msf_constraints')
        prioritize=c4.checkbox('Priorizar tiendas con PRIORITY usando Torton',value=True,key='msf_priority')
        c5,c6=st.columns(2)
        split_orders=c5.checkbox('Partir pedidos',value=True,key='msf_split_orders')
        split_capacity=c6.number_input('CAPACIDAD UNO máxima por parte',min_value=1,max_value=100,value=4,step=1,disabled=not split_orders,key='msf_split_capacity')
    uploaded=st.file_uploader('Archivo de pedidos',type=['xlsx','xlsm','csv'],accept_multiple_files=False,key='msf_excel_inputs',help='Se conservan únicamente las columnas marcadas como DispatchTrack en la fila 1.')
    run=st.button('EJECUTAR PLANEACIÓN →',type='primary',use_container_width=True,disabled=uploaded is None,key='msf_run_excel_shell')
    st.markdown('<div class="msf-note">La configuración se lee desde FLOTA y DIRECCIONES en DATA_LOGISTICS.</div>',unsafe_allow_html=True); st.markdown('</div>',unsafe_allow_html=True)
    if run and uploaded is not None:
        with st.spinner('Leyendo pedidos, FLOTA y asignando unidades…'):
            try:
                headers,rows=_read_dispatchtrack_file(uploaded)
                filtered_fleet=[u for u in fleet_config if u['UNIDAD'] in selected_units]
                filtered_directions={k:dict(v) for k,v in direction_config.items() if not selected_cities or _norm(v.get('CITY','')) in {_norm(city) for city in selected_cities}}
                if not prioritize:
                    for value in filtered_directions.values(): value['PRIORITY']=0
                result,warnings=_assign_and_split(headers,rows,filtered_fleet,filtered_directions,selected_cities,respect_constraints=use_constraints,split_orders=split_orders,split_capacity=split_capacity); csv_bytes=_to_xlsx(headers,result)
            except Exception as exc: st.error(f'No se pudo ejecutar la planeación: {exc}'); return
        st.success(f'Planeación generada: {len(rows):,} pedidos de entrada → {len(result):,} líneas de salida.')
        for warning in warnings: st.info(warning)
        st.session_state['msf_csv_bytes']=csv_bytes; st.session_state['msf_csv_name']=f'MSF_DISPATCHTRACK_{datetime.now():%Y%m%d_%H%M%S}.xlsx'; st.session_state['msf_result_count']=len(result)
    if st.session_state.get('msf_csv_bytes'):
        st.download_button('DESCARGAR XLSX DISPATCHTRACK →',data=st.session_state['msf_csv_bytes'],file_name=st.session_state['msf_csv_name'],mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',use_container_width=True,key='msf_download_shell_csv')
    if st.button('← Volver a Mother Base',key='msf_back_home'): st.session_state['mb_module']='MOTHER BASE'; st.rerun()







