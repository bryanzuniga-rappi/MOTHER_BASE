"""Envío por correo del paquete de una corrida de Les Enfants Terribles.

Posición: acción posterior a la corrida (botón "Enviar por mail" en resultados).
Entrada: el dict ``run`` de execute_planning (archivos en disco) y la lista de destinatarios.
Salida: un solo correo con COPÉRNICO, BASE_TRANSFERS, Sin recomendación y OVERVIEW.
Regla clave: usa la cuenta de Google autorizada por OAuth (mismo cliente que
Militaires) con el permiso ``gmail.send``; un token autorizado solo para Drive
no puede enviar correo y se explica cómo reautorizarlo.
"""

from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Callable, Iterable

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

GMAIL_SEND_SCOPE = "https://www.googleapis.com/auth/gmail.send"
OAUTH_TOKEN_ENDPOINT = "https://oauth2.googleapis.com/token"

DEFAULT_RECIPIENTS = (
    "javier.ballesteros@rappi.com",
    "bryan.zuniga@rappi.com",
    "dorian.santiago@rappi.com",
)

# Nombres de los adjuntos (se arman con la fecha de la corrida, dd-mm-YYYY).
REPORT_NAME = "Reporte_Planeación_{date}.xlsx"
NO_RECOMMENDATION_NAME = "Sin_Recomendación_{date}.csv"
OVERVIEW_NAME = "OVERVIEW_{date}.csv"
COPERNICO_NAME = "Copernico_{warehouse}_{date}.csv"
PACKAGE_ZIP_NAME = "Planeacion_{date}.zip"
DRIVE_FOLDER_ID = "1SfvpuHo0uhZLQb8pG5HZuJFH0bJPMu1f"

# Gmail acepta 25 MB por mensaje ya codificado en base64 (+37 %): el tope de los
# adjuntos crudos es ~18 MB. Arriba de eso se comprimen en un solo ZIP.
MAX_RAW_ATTACHMENT_BYTES = 18 * 1024 * 1024
# Arriba de este tamaño el mensaje se sube por la ruta de carga (media upload).
INLINE_SEND_LIMIT_BYTES = 4 * 1024 * 1024

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv"
ZIP_MIME = "application/zip"

_EMAIL_PATTERN = re.compile(r"^[^@\s,;<>]+@[^@\s,;<>]+\.[^@\s,;<>]+$")


class MailError(Exception):
    """Base de los errores de envío; el mensaje es apto para mostrarse al usuario."""


class MailConfigError(MailError):
    """Faltan secretos de OAuth."""


class MailAuthError(MailError):
    """Google rechazó las credenciales o el permiso."""


class MailPackageError(MailError):
    """No se pudo armar el paquete de archivos."""


@dataclass
class Attachment:
    name: str
    data: bytes
    mime: str

    @property
    def size(self) -> int:
        return len(self.data)


@dataclass
class MailPackage:
    attachments: list[Attachment]
    notes: list[str] = field(default_factory=list)
    zipped: bool = False


@dataclass
class MailResult:
    recipients: list[str]
    subject: str
    attachments: list[tuple[str, int]]
    zipped: bool
    message_id: str
    notes: list[str] = field(default_factory=list)


# --- destinatarios ---------------------------------------------------------------


def parse_recipients(text: str) -> tuple[list[str], list[str]]:
    """(válidos, inválidos) de un texto con correos separados por saltos de línea,
    comas, punto y coma o espacios. Normaliza a minúsculas y quita duplicados."""
    valid: list[str] = []
    invalid: list[str] = []
    for token in re.split(r"[\s,;]+", text or ""):
        token = token.strip().strip("<>").lower()
        if not token:
            continue
        if _EMAIL_PATTERN.match(token):
            if token not in valid:
                valid.append(token)
        elif token not in invalid:
            invalid.append(token)
    return valid, invalid


# --- fecha y nombres ------------------------------------------------------------


def date_label(run: dict[str, Any]) -> str:
    """Fecha de la corrida como dd-mm-YYYY."""
    raw = run.get("run_date")
    if raw:
        return date.fromisoformat(str(raw)).strftime("%d-%m-%Y")
    match = re.search(r"(\d{2}-\d{2}-\d{4})", Path(str(run.get("zip", ""))).name)
    return match.group(1) if match else date.today().strftime("%d-%m-%Y")


def drive_package_name(run: dict[str, Any]) -> str:
    """Nombre del ZIP operativo que se subirá a la carpeta de Drive."""
    origins = [str(int(origin)) for origin in run.get("origins", [])]
    origins_label = "_".join(origins) or "SIN_ORIGEN"
    return f"TR_{origins_label}_{date_label(run)}.zip"


def upload_run_package_to_drive(run: dict[str, Any]) -> str:
    """Arma el ZIP completo y lo sube a la carpeta de Drive operativa."""
    package = build_package(run)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for attachment in package.attachments:
            archive.writestr(attachment.name, attachment.data)
    payload = buffer.getvalue()
    filename = drive_package_name(run)
    try:
        from googleapiclient.http import MediaIoBaseUpload
        from modules.militaires_sans_frontieres import get_drive_service

        service = get_drive_service()
        metadata = {
            "name": filename,
            "parents": [DRIVE_FOLDER_ID],
            "mimeType": ZIP_MIME,
        }
        media = MediaIoBaseUpload(io.BytesIO(payload), mimetype=ZIP_MIME, resumable=True)
        response = service.files().create(
            body=metadata,
            media_body=media,
            supportsAllDrives=True,
            fields="id,name,webViewLink",
        ).execute()
    except Exception as exc:
        raise MailError(f"No se pudo subir el ZIP a Drive: {exc}") from exc
    return str(response.get("webViewLink") or response.get("id") or filename)


def copernico_attachment_names(
    inputs: Iterable[dict[str, Any]], label: str
) -> list[tuple[dict[str, Any], str]]:
    """Copernico_<bodega>_<fecha>.csv por archivo subido. Un archivo con varias
    bodegas lleva todas (444-831); sin bodega detectada conserva su nombre. Si dos
    archivos dan el mismo nombre se agrega un sufijo _2, _3…"""
    used: dict[str, int] = {}
    named: list[tuple[dict[str, Any], str]] = []
    for item in inputs:
        warehouses = sorted(item.get("warehouses") or [])
        if warehouses:
            base = COPERNICO_NAME.format(
                warehouse="-".join(str(w) for w in warehouses), date=label
            )
        else:
            stem = re.sub(r"[^\w.-]+", "_", Path(str(item["path"])).stem).strip("_")
            base = f"Copernico_{stem or 'archivo'}_{label}.csv"
        used[base] = used.get(base, 0) + 1
        name = base if used[base] == 1 else base.replace(".csv", f"_{used[base]}.csv")
        named.append((item, name))
    return named


def preview_names(run: dict[str, Any]) -> list[str]:
    """Nombres de los adjuntos, sin leer ni construir nada (para mostrarlos antes)."""
    label = date_label(run)
    names = [name for _item, name in copernico_attachment_names(run.get("copernico_inputs", []), label)]
    names.append(REPORT_NAME.format(date=label))
    if _find_output(run, "Fountain9_Sin_Recomendacion"):
        names.append(NO_RECOMMENDATION_NAME.format(date=label))
    names.append(OVERVIEW_NAME.format(date=label))
    return names


# --- armado de los archivos ---------------------------------------------------------


def _find_output(run: dict[str, Any], prefix: str) -> Path | None:
    for path in run.get("files", []):
        path = Path(path)
        if path.name.startswith(prefix):
            return path
    return None


def estimated_sheet_bytes(report_path: Path, sheet_name: str = "BASE_TRANSFERS") -> int | None:
    """Tamaño comprimido de una hoja dentro de un .xlsx, leído del directorio del
    ZIP (sin abrir el libro). Sirve para saber, antes de esperar minutos, si la
    hoja cabrá en un correo. None si no se pudo determinar."""
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(report_path) as archive:
            ns_main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            ns_rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            relation_id = next(
                sheet.get(f"{ns_rel}id")
                for sheet in workbook.iter(f"{ns_main}sheet")
                if sheet.get("name") == sheet_name
            )
            rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(
                rel.get("Target") for rel in rels if rel.get("Id") == relation_id
            ).lstrip("/")
            member = target if target.startswith("xl/") else f"xl/{target}"
            return archive.getinfo(member).compress_size
    except Exception:
        return None


def base_transfers_xlsx_bytes(report_path: Path) -> bytes:
    """Un .xlsx con solo la hoja BASE_TRANSFERS del reporte, copiada en streaming
    (no carga el libro completo en memoria)."""
    try:
        source = openpyxl.load_workbook(report_path, read_only=True, data_only=True)
    except Exception as exc:
        raise MailPackageError(f"No pude abrir el reporte de planeación: {exc}") from exc
    try:
        if "BASE_TRANSFERS" not in source.sheetnames:
            raise MailPackageError("El reporte no trae la hoja BASE_TRANSFERS.")
        target = openpyxl.Workbook(write_only=True)
        sheet = target.create_sheet("BASE_TRANSFERS")
        sheet.freeze_panes = "A2"
        rows = source["BASE_TRANSFERS"].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise MailPackageError("La hoja BASE_TRANSFERS está vacía.")
        bold = Font(bold=True)
        header_cells = []
        for value in header:
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = bold
            header_cells.append(cell)
        sheet.append(header_cells)
        for row in rows:
            values = list(row)
            # En modo streaming una celda vacía final necesita estilo para
            # conservar el ancho de la fila al volver a leer el adjunto.
            if values and values[-1] is None:
                last = WriteOnlyCell(sheet, value=None)
                last.font = Font(name="Calibri")
                values[-1] = last
            sheet.append(values)
        buffer = io.BytesIO()
        target.save(buffer)
        return buffer.getvalue()
    finally:
        source.close()


def overview_csv_bytes(run: dict[str, Any]) -> bytes:
    """El mismo "Overview general" que se ve en pantalla, como CSV (con BOM)."""
    from modules.les_enfants_terribles import ordered_breakdown_rows, rows_to_csv_bytes

    rows = ordered_breakdown_rows(run.get("status_counts", {}), run.get("status_swa"))
    if not rows:
        raise MailPackageError("La corrida no tiene overview que enviar.")
    return rows_to_csv_bytes(rows)


def build_package(run: dict[str, Any]) -> MailPackage:
    """Arma los adjuntos de una corrida. Lanza MailPackageError si no hay archivos."""
    if run.get("simulation"):
        raise MailPackageError(
            "La corrida fue en modo simulación: no generó archivos que enviar."
        )
    label = date_label(run)
    notes: list[str] = []
    attachments: list[Attachment] = []

    inputs = run.get("copernico_inputs", [])
    if not inputs:
        notes.append("No se subió COPÉRNICO en esta corrida: el correo va sin esos archivos.")
    for item, name in copernico_attachment_names(inputs, label):
        path = Path(item["path"])
        if not path.exists():
            raise MailPackageError(
                f"El archivo COPÉRNICO {path.name} ya no está disponible; vuelve a ejecutar la planeación."
            )
        attachments.append(Attachment(name, path.read_bytes(), CSV_MIME))

    report = _find_output(run, "Reporte_Planeacion_")
    if report is None or not report.exists():
        raise MailPackageError(
            "No encontré el reporte de planeación de esta corrida; vuelve a ejecutarla."
        )
    # Falla rápido: leer y reescribir un reporte enorme tarda minutos y de todos
    # modos no cabría en un correo.
    estimate = estimated_sheet_bytes(report)
    if estimate is not None and estimate > MAX_RAW_ATTACHMENT_BYTES:
        raise MailPackageError(
            f"La hoja BASE_TRANSFERS pesa unos {estimate / 1024 / 1024:,.1f} MB y Gmail acepta "
            f"unos {MAX_RAW_ATTACHMENT_BYTES / 1024 / 1024:,.0f} MB por correo. Descárgala del "
            "ZIP de la corrida y compártela por Drive."
        )
    attachments.append(
        Attachment(REPORT_NAME.format(date=label), base_transfers_xlsx_bytes(report), XLSX_MIME)
    )

    no_recommendation = _find_output(run, "Fountain9_Sin_Recomendacion")
    if no_recommendation is not None and no_recommendation.exists():
        attachments.append(
            Attachment(
                NO_RECOMMENDATION_NAME.format(date=label),
                no_recommendation.read_bytes(),
                CSV_MIME,
            )
        )
    else:
        notes.append("Esta corrida no generó el archivo de Fountain9 sin recomendación.")

    attachments.append(
        Attachment(OVERVIEW_NAME.format(date=label), overview_csv_bytes(run), CSV_MIME)
    )
    return _enforce_size_limit(MailPackage(attachments, notes), label)


def _enforce_size_limit(package: MailPackage, label: str) -> MailPackage:
    total = sum(a.size for a in package.attachments)
    if total <= MAX_RAW_ATTACHMENT_BYTES:
        return package
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for attachment in package.attachments:
            archive.writestr(attachment.name, attachment.data)
    zipped = buffer.getvalue()
    if len(zipped) > MAX_RAW_ATTACHMENT_BYTES:
        raise MailPackageError(
            f"Los archivos pesan {total / 1024 / 1024:,.1f} MB ({len(zipped) / 1024 / 1024:,.1f} MB "
            f"comprimidos) y Gmail acepta unos {MAX_RAW_ATTACHMENT_BYTES / 1024 / 1024:,.0f} MB por correo."
        )
    return MailPackage(
        [Attachment(PACKAGE_ZIP_NAME.format(date=label), zipped, ZIP_MIME)],
        package.notes
        + [
            f"Los archivos pesaban {total / 1024 / 1024:,.1f} MB, así que se enviaron "
            "comprimidos en un solo ZIP."
        ],
        zipped=True,
    )


# --- mensaje -------------------------------------------------------------------------


def build_subject(run: dict[str, Any]) -> str:
    return f"Les Enfants Terribles · Planeación {date_label(run)}"


def build_body(run: dict[str, Any], package: MailPackage) -> str:
    origins = ", ".join(str(o) for o in run.get("origins", [])) or "—"
    lines = [
        f"Planeación de Les Enfants Terribles del {date_label(run)}.",
        "",
        f"Orígenes: {origins}",
        f"Unidades planeadas: {int(run.get('units', 0) or 0):,}",
        f"Tareas: {int(run.get('tasks', 0) or 0):,}",
    ]
    swa = run.get("swa_report") or {}
    if swa.get("enabled"):
        lines.append(f"SWA ganado: {float(swa.get('swa_ganado', 0.0)):,.4f}")
    lines += ["", "Archivos adjuntos:"]
    lines += [f"- {a.name}" for a in package.attachments]
    lines += [f"- Nota: {note}" for note in package.notes]
    lines += ["", f"Versión del motor: {run.get('build', '—')}", "Enviado desde Mother Base."]
    return "\n".join(lines)


def compose_message(
    recipients: list[str], subject: str, body: str, attachments: list[Attachment]
) -> EmailMessage:
    message = EmailMessage()
    message["To"] = ", ".join(recipients)
    message["Subject"] = subject
    message.set_content(body)
    for attachment in attachments:
        maintype, _, subtype = attachment.mime.partition("/")
        message.add_attachment(
            attachment.data, maintype=maintype, subtype=subtype, filename=attachment.name
        )
    return message


# --- Gmail ---------------------------------------------------------------------------

REAUTH_MESSAGE = (
    "El token de Google guardado en Secrets solo tiene permiso de Drive y no puede "
    "enviar correo. Hace falta autorizarlo una vez más con el permiso de Gmail: en "
    "Militaires Sans Frontières abre «Conectar mi cuenta personal de Google (OAuth)», "
    "autoriza (ahora pide Drive y enviar correo) y pega el nuevo refresh token en "
    "Secrets (DATA_DASHBOARD_OAUTH_REFRESH_TOKEN, o MAIL_OAUTH_REFRESH_TOKEN si lo "
    "quieres aparte). Además, la API de Gmail debe estar habilitada en el proyecto de "
    "Google Cloud del cliente OAuth."
)


def _read_streamlit_secrets() -> dict[str, str]:
    keys = (
        "DATA_DASHBOARD_OAUTH_CLIENT_ID",
        "DATA_DASHBOARD_OAUTH_CLIENT_SECRET",
        "DATA_DASHBOARD_OAUTH_REFRESH_TOKEN",
        "MAIL_OAUTH_REFRESH_TOKEN",
    )
    try:
        import streamlit as st

        return {key: str(st.secrets.get(key, "") or "") for key in keys}
    except Exception:
        return {key: "" for key in keys}


def build_gmail_service(
    secrets: dict[str, str] | None = None,
    credentials_factory: Callable[..., Any] | None = None,
    service_builder: Callable[..., Any] | None = None,
) -> Any:
    """Cliente de la API de Gmail con el OAuth de usuario ya configurado.

    ``secrets``, ``credentials_factory`` y ``service_builder`` existen para
    probarlo sin red; en producción se leen de st.secrets y de las librerías de Google."""
    secrets = secrets if secrets is not None else _read_streamlit_secrets()
    client_id = secrets.get("DATA_DASHBOARD_OAUTH_CLIENT_ID", "")
    client_secret = secrets.get("DATA_DASHBOARD_OAUTH_CLIENT_SECRET", "")
    refresh_token = secrets.get("MAIL_OAUTH_REFRESH_TOKEN", "") or secrets.get(
        "DATA_DASHBOARD_OAUTH_REFRESH_TOKEN", ""
    )
    missing = [
        name
        for name, value in (
            ("DATA_DASHBOARD_OAUTH_CLIENT_ID", client_id),
            ("DATA_DASHBOARD_OAUTH_CLIENT_SECRET", client_secret),
            ("DATA_DASHBOARD_OAUTH_REFRESH_TOKEN (o MAIL_OAUTH_REFRESH_TOKEN)", refresh_token),
        )
        if not value
    ]
    if missing:
        raise MailConfigError(
            "Faltan secretos para enviar correo: " + ", ".join(missing) + "."
        )

    if credentials_factory is None:
        from google.oauth2.credentials import Credentials

        credentials_factory = Credentials
    credentials = credentials_factory(
        None,
        refresh_token=refresh_token,
        client_id=client_id,
        client_secret=client_secret,
        token_uri=OAUTH_TOKEN_ENDPOINT,
        scopes=[GMAIL_SEND_SCOPE],
    )
    try:
        if service_builder is None:
            from google.auth.transport.requests import Request

            credentials.refresh(Request())
        else:
            credentials.refresh(None)
    except Exception as exc:
        text = str(exc).lower()
        if "invalid_scope" in text or "unauthorized_client" in text:
            raise MailAuthError(REAUTH_MESSAGE) from exc
        if "invalid_grant" in text:
            raise MailAuthError(
                "Google rechazó el refresh token (invalid_grant): venció o fue revocado. "
                "Vuelve a autorizar la cuenta y actualiza el token en Secrets."
            ) from exc
        raise MailAuthError(f"No pude autenticarme con Google: {exc}") from exc

    if service_builder is None:
        from googleapiclient.discovery import build

        service_builder = build
    return service_builder("gmail", "v1", credentials=credentials, cache_discovery=False)


def friendly_send_error(exc: Exception) -> MailError:
    """Traduce un error de la API de Gmail a un mensaje accionable."""
    text = str(exc)
    lowered = text.lower()
    if "insufficient authentication scopes" in lowered or "insufficientpermissions" in lowered:
        return MailAuthError(REAUTH_MESSAGE)
    if "accessnotconfigured" in lowered or "has not been used in project" in lowered:
        return MailConfigError(
            "La API de Gmail no está habilitada en el proyecto de Google Cloud del cliente "
            "OAuth. Habilítala en «APIs y servicios» y vuelve a intentar."
        )
    if "ratelimit" in lowered or "quota" in lowered:
        return MailError("Google limitó los envíos por ahora (cuota). Intenta de nuevo en unos minutos.")
    if "invalid to header" in lowered or "invalid_to" in lowered:
        return MailError("Gmail rechazó uno de los destinatarios. Revisa las direcciones.")
    return MailError(f"Gmail no pudo enviar el correo: {text}")


def send_message(
    service: Any,
    message: EmailMessage,
    media_factory: Callable[[bytes], Any] | None = None,
) -> str:
    """Envía el mensaje con la cuenta autorizada ("me"); devuelve el id del mensaje."""
    raw_bytes = message.as_bytes()
    try:
        messages = service.users().messages()
        if len(raw_bytes) > INLINE_SEND_LIMIT_BYTES:
            if media_factory is None:
                from googleapiclient.http import MediaIoBaseUpload

                def media_factory(data: bytes) -> Any:  # noqa: E306
                    return MediaIoBaseUpload(
                        io.BytesIO(data), mimetype="message/rfc822", resumable=True
                    )

            response = messages.send(userId="me", media_body=media_factory(raw_bytes)).execute()
        else:
            raw = base64.urlsafe_b64encode(raw_bytes).decode("ascii")
            response = messages.send(userId="me", body={"raw": raw}).execute()
    except MailError:
        raise
    except Exception as exc:
        raise friendly_send_error(exc) from exc
    return str((response or {}).get("id", ""))


def send_run_package(
    run: dict[str, Any],
    recipients: list[str],
    service: Any | None = None,
    media_factory: Callable[[bytes], Any] | None = None,
) -> MailResult:
    """Arma el paquete de la corrida y lo envía en un solo correo."""
    if not recipients:
        raise MailError("Agrega al menos un destinatario.")
    package = build_package(run)          # primero los archivos: falla antes de autenticar
    message = compose_message(
        recipients, build_subject(run), build_body(run, package), package.attachments
    )
    service = service or build_gmail_service()
    message_id = send_message(service, message, media_factory)
    return MailResult(
        recipients=list(recipients),
        subject=build_subject(run),
        attachments=[(a.name, a.size) for a in package.attachments],
        zipped=package.zipped,
        message_id=message_id,
        notes=package.notes,
    )
