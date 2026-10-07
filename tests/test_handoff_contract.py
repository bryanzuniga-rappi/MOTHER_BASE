"""Regresiones del contrato DOI literal, ledger residual y entregables."""
import csv
import io
from datetime import date
from pathlib import Path

import openpyxl
import pytest

import modelo_abasto as engine
import modules.les_enfants_terribles as m
from tests.test_fountain9_comparison import make_catalogs
from tests.e2e_fixture import FakeUpload, build_workbook_bytes


def catalogs(**overrides):
    values = dict(
        stock_base={(444, 10): 20, (811, 10): 20},
        volume_m3={10: 1}, store_capacity={100: 100, 200: 100},
        stores={i: {"city": "CDMX", "city_norm": "CDMX", "warehouse_name": str(i)}
                for i in (444, 811, 425, 856, 100, 200)},
    )
    values.update(overrides)
    return make_catalogs(**values)


def row(qty=20, source=444, destination=100, **extra):
    return dict(WAREHOUSE_DESTINATION=destination, RETAIL_ID=10,
                CURRENT_INVENTORY=0, PREDICTED_DEMAND=qty,
                PREDICTED_OPENING_INVENTORY=0, MOV_ORIGINAL=qty,
                INPUT_ROW=1, SKU_NAME="SKU", FORCED_TARGET=qty,
                FORCED_SOURCE=source, FORCED_RULE="FOUNTAIN9_DOI",
                ALLOW_PARTIAL=True, **extra)


@pytest.mark.parametrize("change,expected,cause", [
    ({"stock_base": {(444, 10): 12}}, 12, "STOCK"),
    ({"copernico_unusable_by_warehouse": {(444, 10): 8}}, 12, "COPÉRNICO"),
    ({"unavailable_stock": {(444, 10): 8}}, 12, "STOCK"),
    ({"store_capacity": {100: 12}}, 12, "TIENDA"),
    ({"rackeados_444": {10}}, 0, "STOCK"),
    ({"route_cost_blocks": {(100, 10)}}, 0, "RUTA"),
    ({"schedule_block_enabled": True, "run_weekday_norm": "LUNES",
      "schedule_days": {(100, 444): frozenset({"MARTES"})}}, 0, "FRECUENCIA"),
])
def test_literal_source_respects_all_limits(change, expected, cause):
    result = engine.plan_transfers([row()], catalogs(**change), engine.Config(origin_warehouses=(444, 811)))
    assert sum(r["QUANTITY"] for r in result.allocation_rows) == expected
    assert all(r["WAREHOUSE_SOURCE"] == 444 for r in result.allocation_rows)
    assert cause in result.base_rows[0]["TIPO_DE_CORTE"]
    assert result.base_rows[0]["CANTIDAD_FALTANTE"] == 20 - expected


@pytest.mark.parametrize("source,cause", [(None, "SIN ORIGEN"), (999, "INVALIDO"), ("abc", "INVALIDO")])
def test_invalid_origin_never_falls_back(source, cause):
    result = engine.plan_transfers([row(source=source)], catalogs(), engine.Config(origin_warehouses=(444, 811)))
    assert not result.allocation_rows
    assert cause in result.base_rows[0]["TIPO_DE_CORTE"]


def test_residual_reuses_task_and_capacity_without_double_consumption():
    cat = catalogs(store_capacity={100: 10})
    config = engine.Config(origin_warehouses=(444,), max_tasks=1)
    naked = engine.plan_transfers([row(3)], cat, config)
    residual = row(7)
    residual.pop("FORCED_SOURCE")
    result = engine.plan_transfers([residual], cat, config, initial_result=naked)
    assert result.tasks_used == 1
    assert sum(r["QUANTITY"] for r in result.allocation_rows) == 10
    assert result.capacity_rows[0]["M3_CONTABILIZADO_CAPACIDAD"] == 10
    assert result.base_rows[-1]["STOCK_REMANENTE_444"] == 10
    assert result.base_rows[-1]["TAREAS_GENERADAS"] == 0


def test_tasks_cut_a_second_instruction_but_not_an_existing_route():
    config = engine.Config(origin_warehouses=(444,), max_tasks=1)
    result = engine.plan_transfers([row(3), row(3, destination=200)], catalogs(), config)
    assert sum(r["QUANTITY"] for r in result.allocation_rows) == 3
    assert "TAREAS" in result.base_rows[1]["TIPO_DE_CORTE"]


@pytest.mark.parametrize("source", [425, 856])
def test_owner_is_reserved_before_otacon_and_task_limit(source):
    cat = catalogs(stock_base={(source, 10): 20},
                   owner_stock={(source, 10, "A"): 12, (source, 10, "B"): 8})
    config = engine.Config(origin_warehouses=(source,), max_tasks=1)
    naked = engine.plan_transfers([row(source=source)], cat, config)
    assert naked.base_rows[0]["CANTIDAD_ASIGNADA"] == 12
    assert "TAREAS" in naked.base_rows[0]["TIPO_DE_CORTE"]
    residual = row(8, source)
    residual.pop("FORCED_SOURCE")
    result = engine.plan_transfers([residual], cat, config, initial_result=naked)
    engine.apply_owner_inventory_partition(result, cat, config)
    assert sum(r["QUANTITY"] for r in result.allocation_rows) == 12
    assert result.tasks_used == 1
    assert {r["OWNER_NAME"] for r in result.allocation_rows} == {"A"}


def csv_plan(rows):
    columns = ["Warehouseid", "SKU ID", "Current Inventory", "Predicted Demand for selected duration",
               "Predicted Opening Inventory", "Replenishment Quantity for Plan Duration (MOV)",
               "Net Inter-Store Transfers", "Allocation (DOI Based)", "Source Id Before Multi Source",
               "Primary Source Id", "Allocated Qty Before Multi Source", "Source Current Inv Before Multi Source",
               "Is Secondary Plan", "Link Type", "DOI", "DOH"]
    out = io.StringIO()
    writer = csv.DictWriter(out, columns)
    writer.writeheader()
    for doi, source, mov in rows:
        writer.writerow(dict(zip(columns, [100, 10, 0, mov, 0, mov, 0, doi, source, 999, 0, 10000, "Yes", "Transfer", 1, 0])))
    return out.getvalue().encode()


def execute(rows, stock=20, **overrides):
    args = dict(uploaded_copernico=[], uploaded_plans=[FakeUpload("f9.csv", csv_plan(rows))],
                database_bytes=build_workbook_bytes({100: "CDMX"}, [10], origin_stock=stock),
                origins=(444,), max_tasks=100, run_date=date(2026, 10, 6), include_insumos=False,
                block_fruver_811=False, block_off_schedule_shipments=False,
                enable_closed_stores_rule=False, minimum_positive_quantity=4)
    args.update(overrides)
    return m.execute_planning(**args)


@pytest.mark.parametrize("doi,mov,final", [(3, 10, 10), (12, 3, 12)])
def test_maximum_keeps_doi_when_optional_columns_off(doi, mov, final):
    run = execute([(doi, 444, mov)], use_extra_mov_columns=False)
    assert run["units"] == final
    assert run["fountain9_audit"][0]["NAKED_EJECUTADO"] == doi
    assert run["tasks"] == 1


def test_duplicates_sum_and_exports_reconcile(tmp_path):
    run = execute([(3, 444, 2), (9, 444, 2)])
    assert run["units"] == 12
    audit = run["fountain9_audit"]
    assert len(audit) == 1 and audit[0]["DOI_SOLICITADO"] == 12
    assert audit[0]["NAKED_EJECUTADO"] == 12
    assert audit[0]["Is Secondary Plan"] == "Yes"
    report = next(Path(p) for p in run["files"] if str(p).endswith(".xlsx"))
    wb = openpyxl.load_workbook(report, read_only=True)
    assert "AUDITORIA_FOUNTAIN9" in wb.sheetnames
    assert "PLANNING_REASON" in next(wb["DETALLE_ASIGNACION"].values)
    wb.close()
    bulk = next(Path(p) for p in run["files"] if Path(p).name == "BulkCD_444.csv")
    with bulk.open(encoding="utf-8-sig") as f:
        assert len(list(csv.DictReader(f))) == 1


@pytest.mark.parametrize("source", ["", 999, "invalid"])
def test_otacon_recovers_missing_or_invalid_f9_origin(source):
    run = execute([(10, source, 5)])
    line = run["fountain9_audit"][0]
    assert line["NAKED_EJECUTADO"] == 0
    assert line["OTACON_RECUPERADO"] == 10
    assert line["FALTANTE_FINAL"] == 0
    assert "ORIGEN" in line["MOTIVO_CORTE_NAKED"]


def test_excluded_store_is_not_reconstructed_by_naked():
    run = execute([(10, 444, 10)], excluded_store_ids=(100,))
    assert run["units"] == 0
    assert run["fountain9_audit"][0]["MOTIVO_CORTE_NAKED"] == "TIENDA_EXCLUIDA"


def test_recovery_is_not_duplicated_for_two_f9_origins():
    run = execute([(8, "", 5), (8, 999, 5)], stock=10)
    audit = run["fountain9_audit"]
    assert sum(r["DOI_SOLICITADO"] for r in audit) == 16
    assert sum(r["OTACON_RECUPERADO"] for r in audit) == 10
    assert sum(r["FALTANTE_FINAL"] for r in audit) == 6


def test_priority_is_bucket_then_data_priority_then_doi_doh():
    config = engine.Config(origin_warehouses=(444,), max_tasks=1)
    result = engine.plan_transfers([row(4, DOI=8, DOH=1), row(4, destination=200, DOI=1, DOH=4)], catalogs(), config)
    assert result.allocation_rows[0]["WAREHOUSE_DESTINATION"] == 200


def test_capacity_rows_survive_empty_residual():
    config = engine.Config(origin_warehouses=(444,))
    initial = engine.plan_transfers([row(4)], catalogs(), config)
    result = engine.plan_transfers([], catalogs(), config, initial_result=initial)
    assert result.capacity_rows[0]["M3_CONTABILIZADO_CAPACIDAD"] == 4


@pytest.mark.parametrize("source", [425, 856])
def test_owner_exports_two_files_with_one_route_each(source):
    database = build_workbook_bytes({100: "CDMX"}, [10], origin=source, origin_stock=20,
                                   extra_rows={"OWNER": [(source, 10, "A", 12), (source, 10, "B", 8)]})
    run = execute([(15, source, 20)], origins=(source,), database_bytes=database)
    assert run["units"] == 20
    assert run["tasks"] == 2
    files = [Path(p) for p in run["files"] if Path(p).name.startswith(f"BulkCD_{source}_")]
    assert len(files) == 2
    quantities = []
    for file in files:
        with file.open(encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 1
        quantities.append(int(rows[0]["QUANTITY"]))
    assert sorted(quantities) == [8, 12]
    assert run["fountain9_audit"][0]["NAKED_EJECUTADO"] == 15


@pytest.mark.parametrize("restriction,cause", [("TIENDAS_CERRADAS", "TIENDA_CERRADA"), ("BLOQUEOS", "SKU_BLOQUEADO")])
def test_closed_and_blocked_instructions_remain_in_audit(restriction, cause):
    database = build_workbook_bytes({100: "CDMX"}, [10], extra_rows={restriction: [(100 if restriction == "TIENDAS_CERRADAS" else 10,)]})
    run = execute([(10, 444, 10)], database_bytes=database, enable_closed_stores_rule=True)
    assert run["units"] == 0
    assert run["fountain9_audit"][0]["MOTIVO_CORTE_NAKED"] == cause


def test_city_block_is_audited_and_never_executed():
    run = execute([(10, 444, 10)], blocked_cities=("CDMX",))
    assert run["units"] == 0
    assert run["fountain9_audit"][0]["MOTIVO_CORTE_NAKED"] == "CIUDAD_BLOQUEADA"


def test_shalashaska_uses_route_activated_by_solidus():
    database = build_workbook_bytes({100: "CDMX"}, [10], origin_stock=30,
        extra_rows={"POR_MERMAR": [(444, 10, 10, 10, "2026-10-01", "2026-10-08")],
                    "SHARE_VENTAS": [(100, 1.0)]})
    run = execute([(0, "", 0)], database_bytes=database, cover_fountain9_hardcodes=False,
                  include_avl_fill=True, include_shalashaska_engine=True)
    assert run["avl"]["units_added"] > 0
    assert run["shalashaska"]["units_evacuated"] > 0
    summary = {r["ENGINE"]: r["UNIDADES"] for r in run["engine_summary_rows"] if r["COBERTURA"] in ("Total", m.NO_COVERAGE)}
    assert summary["Naked"] == 0
    assert summary["Shalashaska"] == run["shalashaska"]["units_evacuated"]


def test_otacon_recovers_from_different_real_origin():
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook_bytes({100: "CDMX"}, [10], origin_stock=3)))
    wb["STOCK"].append([811, 10, 7, 0])
    wb["TIENDA"].append(["CDMX", 811, "O811"])
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    run = execute([(10, 444, 10)], database_bytes=out.getvalue(), origins=(444, 811))
    audit = run["fountain9_audit"][0]
    assert (audit["NAKED_EJECUTADO"], audit["OTACON_RECUPERADO"], audit["FALTANTE_FINAL"]) == (3, 7, 0)
    assert sum(r["TAREAS"] for r in run["engine_summary_rows"] if r["COBERTURA"] in ("Total", m.NO_COVERAGE)) == run["tasks"] == 2
    assert run["analytics"]["summary"]["naked_target_units"] == 10
    assert run["analytics"]["summary"]["target_units"] == 10
    report = next(Path(p) for p in run["files"] if str(p).endswith(".xlsx"))
    wb = openpyxl.load_workbook(report, read_only=True, data_only=True)
    metrics = {r[0]: r[1] for r in wb["RESUMEN"].iter_rows(values_only=True) if len(r) >= 2}
    assert metrics["UNIDADES_OBJETIVO"] == metrics["UNIDADES_ASIGNADAS"] == 10
    assert metrics["UNIDADES_FALTANTES"] == 0
    wb.close()


def test_pdf_and_ui_include_doi_audit():
    from pypdf import PdfReader
    from tests.test_engine_report_layout import _render
    run = execute([(3, 444, 10)])
    pdf = next(Path(p) for p in run["files"] if str(p).endswith(".pdf"))
    text = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
    assert "Naked ejecutado literalmente" in text
    assert "Otacon recuperado" in text
    assert any("Auditoría Fountain9 DOI" in call for call in _render(run) if isinstance(call, str))


def test_randomized_shared_ledger_invariants():
    from collections import Counter
    import random
    rng = random.Random(42)
    for _ in range(40):
        stock = rng.randint(0, 30)
        capacity = rng.randint(0, 25)
        budget = rng.randint(0, 2)
        cat = catalogs(stock_base={(444, 10): stock}, store_capacity={100: capacity, 200: capacity})
        config = engine.Config(origin_warehouses=(444,), max_tasks=budget)
        initial = engine.plan_transfers([row(rng.randint(1, 20), destination=d) for d in (100, 200)], cat, config)
        residuals = [row(rng.randint(1, 20), destination=d) for d in (100, 200)]
        for r in residuals:
            r.pop("FORCED_SOURCE")
        result = engine.plan_transfers(residuals, cat, config, initial_result=initial)
        assert sum(r["QUANTITY"] for r in result.allocation_rows) <= stock
        by_store = Counter()
        for r in result.allocation_rows:
            by_store[r["WAREHOUSE_DESTINATION"]] += r["QUANTITY"]
        assert all(q <= capacity for q in by_store.values())
        assert result.tasks_used == len(by_store) <= budget
