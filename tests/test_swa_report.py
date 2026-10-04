"""Pruebas de la hoja SWA (carga en Catalogs) y del reporte central
build_swa_report: SWA ganado/perdido sobre todo el universo de quiebres
del catálogo."""

from types import SimpleNamespace

import modelo_abasto as engine
from tests._streamlit_stub import install as _install_streamlit_stub

_install_streamlit_stub()

import modules.les_enfants_terribles as m  # noqa: E402


def make_catalogs(**overrides) -> engine.Catalogs:
    base = dict(
        volume_m3={}, blocked_products=set(), route_cost_blocks=set(),
        store_priority={}, high_value={}, rackeados_444=set(),
        store_capacity={}, copernico_unusable_444={}, unavailable_stock={},
        stock_base={}, golden_infaltables=set(), stores={}, storage={},
        warnings=[],
    )
    base.update(overrides)
    return engine.Catalogs(**base)


def _catalog_row(destination, sku):
    return {"WAREHOUSE_DESTINATION": destination, "RETAIL_ID": sku, "ADU": 1.0}


def _alloc_row(destination, sku, quantity):
    return {
        "WAREHOUSE_DESTINATION": destination,
        "RETAIL_ID": sku,
        "QUANTITY": quantity,
    }


# --- build_swa_report -------------------------------------------------

def test_disabled_when_no_swa_data_loaded():
    catalogs = make_catalogs(stock_base={(100, 10): 0})
    result = SimpleNamespace(allocation_rows=[])
    report = m.build_swa_report([_catalog_row(100, 10)], catalogs, result)
    assert report["enabled"] is False


def test_ganado_is_binary_any_positive_shipment_captures_full_value():
    """No es proporcional. Enviar 1 unidad de las 50 necesarias captura el SWA completo
    igual que enviar las 50.
    """
    catalogs = make_catalogs(
        stock_base={(100, 10): 0},
        swa_potential_gain={(100, 10): 2.5},
    )
    result = SimpleNamespace(allocation_rows=[_alloc_row(100, 10, 1)])
    report = m.build_swa_report([_catalog_row(100, 10)], catalogs, result)
    assert report["swa_ganado"] == 2.5
    assert report["casos_ganados"] == 1


def test_perdido_when_stockout_stays_uncovered():
    catalogs = make_catalogs(
        stock_base={(100, 10): 0},
        swa_potential_gain={(100, 10): 3.1},
    )
    result = SimpleNamespace(allocation_rows=[])  # nadie lo cubrió
    report = m.build_swa_report([_catalog_row(100, 10)], catalogs, result)
    assert report["swa_perdido"] == 3.1
    assert report["swa_ganado"] == 0.0
    assert report["casos_perdidos"] == 1


def test_universo_incluye_quiebres_sin_intento_de_cobertura():
    """El universo es TODO quiebre del catálogo, aunque ningún engine haya intentado
    cubrirlo (sin stock en ningún origen, por ejemplo) — debe seguir contando como
    'perdido'.
    """
    catalogs = make_catalogs(
        stock_base={(100, 10): 0, (200, 20): 0},
        swa_potential_gain={(100, 10): 1.0, (200, 20): 2.0},
    )
    result = SimpleNamespace(allocation_rows=[])
    report = m.build_swa_report(
        [_catalog_row(100, 10), _catalog_row(200, 20)], catalogs, result
    )
    assert report["universo_total"] == 2
    assert report["casos_perdidos"] == 2
    assert report["swa_perdido"] == 3.0


def test_sku_con_stock_positivo_no_cuenta_en_el_universo():
    """No es un quiebre -> SWA no aplica, ni suma a ganado ni a perdido."""
    catalogs = make_catalogs(
        stock_base={(100, 10): 50},
        swa_potential_gain={(100, 10): 9.9},
    )
    result = SimpleNamespace(allocation_rows=[])
    report = m.build_swa_report([_catalog_row(100, 10)], catalogs, result)
    assert report["universo_total"] == 0
    assert report["swa_ganado"] == 0.0
    assert report["swa_perdido"] == 0.0


def test_sku_ausente_de_hoja_swa_usa_cero_por_descarte():
    """Si no está en la hoja SWA, es 0 — nunca bloquea el reporte, el caso se sigue
    contando (ganado o perdido) solo que con SWA=0.
    """
    catalogs = make_catalogs(
        stock_base={(100, 10): 0},
        swa_potential_gain={},  # SKU 10 no aparece aquí
    )
    result = SimpleNamespace(allocation_rows=[_alloc_row(100, 10, 5)])
    report = m.build_swa_report([_catalog_row(100, 10)], catalogs, result)
    assert report["casos_ganados"] == 1
    assert report["swa_ganado"] == 0.0
    assert report["casos_sin_swa_registrado"] == 1


def test_multiple_engines_same_key_counted_once_not_double():
    """Si AVL y Liquid ambos mandan algo a la misma tienda-SKU, el SWA se
    captura UNA vez, no se duplica por tener dos líneas."""
    catalogs = make_catalogs(
        stock_base={(100, 10): 0},
        swa_potential_gain={(100, 10): 4.0},
    )
    result = SimpleNamespace(
        allocation_rows=[_alloc_row(100, 10, 2), _alloc_row(100, 10, 3)]
    )
    report = m.build_swa_report([_catalog_row(100, 10)], catalogs, result)
    assert report["swa_ganado"] == 4.0  # no 8.0
    assert report["casos_ganados"] == 1


def test_duplicate_catalog_rows_deduplicated():
    """Si catalog_rows trae la misma combinación dos veces, no debe
    contarse dos veces en el universo."""
    catalogs = make_catalogs(
        stock_base={(100, 10): 0},
        swa_potential_gain={(100, 10): 1.5},
    )
    result = SimpleNamespace(allocation_rows=[])
    report = m.build_swa_report(
        [_catalog_row(100, 10), _catalog_row(100, 10)], catalogs, result
    )
    assert report["universo_total"] == 1


def test_insumos_only_coverage_counts_as_ganado():
    """Una tienda-SKU cubierta ÚNICAMENTE por Insumos (que no vive en
    result.allocation_rows) debe contar como 'ganado', no 'perdido'.
    """
    catalogs = make_catalogs(
        stock_base={(100, 10): 0},
        swa_potential_gain={(100, 10): 1.8},
    )
    result = SimpleNamespace(allocation_rows=[])  # nada por la vía regular
    insumos_summary = {"assigned_keys": {(100, 10)}}
    report = m.build_swa_report(
        [_catalog_row(100, 10)], catalogs, result, insumos_summary
    )
    assert report["swa_ganado"] == 1.8
    assert report["casos_ganados"] == 1
    assert report["swa_perdido"] == 0.0


def test_swa_report_works_without_insumos_summary():
    """Insumos_summary es opcional (None por default) — no debe tronar."""
    catalogs = make_catalogs(
        stock_base={(100, 10): 0}, swa_potential_gain={(100, 10): 1.0}
    )
    result = SimpleNamespace(allocation_rows=[])
    report = m.build_swa_report([_catalog_row(100, 10)], catalogs, result)
    assert report["swa_perdido"] == 1.0


# --- carga de la hoja SWA en Catalogs (modelo_abasto.py) -----------------

def test_iter_swa_records_absent_sheet_yields_nothing():
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("OTRA_HOJA")
    assert list(engine.iter_swa_records(wb)) == []


def test_iter_swa_records_reads_real_header_position(tmp_path):
    """La hoja SWA real trae metadata de Aleph antes del encabezado (no
    está en la fila 1) — confirma que el scan de 40 filas lo encuentra."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("SWA")
    ws.append([None, "Aleph | Explore"])
    ws.append([None, "Last Updated", "October 01, 2026, 12:00:00 PM CST"])
    for _ in range(4):
        ws.append([])
    ws.append(["WAREHOUSE_ID", "PRODUCT_ID", "SWA_POTENTIAL_GAIN_COUNTRY"])
    ws.append([100, 10, 0.0283])
    path = tmp_path / "swa_test.xlsx"
    wb.save(path)

    wb2 = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = list(engine.iter_swa_records(wb2))
    assert len(records) == 1
    assert records[0]["WAREHOUSE_ID"] == 100
    assert records[0]["PRODUCT_ID"] == 10
    assert records[0]["SWA_POTENTIAL_GAIN_COUNTRY"] == 0.0283


def test_swa_potential_gain_keeps_max_on_duplicate_key(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("SWA")
    ws.append(["WAREHOUSE_ID", "PRODUCT_ID", "SWA_POTENTIAL_GAIN_COUNTRY"])
    ws.append([100, 10, 1.5])
    ws.append([100, 10, 3.2])
    for name, headers, extra_row in [
        ("VOLUMETRIA", ["SKU", "PALLETS"], None),
        ("BLOQUEOS_FORANEAS", ["SKU"], None),
        ("BLOQUEOS", ["PRODUCT_ID"], None),
        ("RUTA_COSTOS", ["Destination", "Catalog ID"], None),
        ("PRIORIDAD", ["WAREHOUSE_ID", "PRIORIDAD"], None),
        ("444_HV", ["EAN", "Category"], None),
        ("831_HV", ["EAN", "Category"], None),
        ("RACKEADOS", ["WHS", "SYNC"], None),
        ("CAP_RECIBO", ["WH_ID", "CAP"], None),
        ("CATALOGO", ["WAREHOUSE_ID", "PRODUCT_ID", "ADU"], None),
        ("KVI", ["WAREHOUSE_ID", "PRODUCT_ID", "KVI"], None),
        ("SHARE_VENTAS", ["WAREHOUSE_ID", "SHARE"], None),
        ("NO_DISPONIBLE", ["WAREHOUSE_ID", "PRODUCT_ID", "STOCK"], None),
        ("POR_MERMAR", ["WAREHOUSE_ID", "PRODUCT_ID", "STOCK_AVAILABLE", "VALUE_STOCK", "ARRIVAL_DATE", "EXPIRATION_DATE"], None),
        ("STOCK", ["WAREHOUSE_ID", "PRODUCT_ID", "STOCK_DISPONIBLE_FINAL"], None),
        ("OWNER", ["WAREHOUSE_ID", "PRODUCT_ID", "OWNER_NAME", "STOCK_DISPONIBLE_FINAL"], None),
        ("INSUMOS", ["WAREHOUSE_DESTINATION", "WAREHOUSE_SOURCE", "RETAIL_ID", "QUANTITY", "PLANNED_DATE", "ROUTE", "DELIVERY_PRIORITY"], None),
        ("GOLDEN_INFALTABLES_ANCHOR", ["WAREHOUSE_ID", "PRODUCT_ID_SYNC", "IS_INFALTABLE", "IS_GOLDEN", "IS_ANCHOR"], None),
        ("TIENDA", ["CITY", "WAREHOUSE_ID", "WAREHOUSE_NAME"], ["Ciudad de México", 100, "STORE"]),
        ("STORAGE", ["PRODUCT_ID", "STORAGE_NAME"], None),
        ("TIENDAS_CERRADAS", ["WAREHOUSE_ID"], None),
        ("SCHEDULE", ["CITY", "WAREHOUSE_ID", "WAREHOUSE_NAME", "ORIGEN", "DAYS"], None),
    ]:
        extra = wb.create_sheet(name)
        extra.append(headers)
        if extra_row:
            extra.append(extra_row)
        if name == "TIENDA":
            extra.append(["Ciudad de México", 444, "O444"])

    path = tmp_path / "data_transfers.xlsx"
    wb.save(path)
    config = engine.Config(origin_warehouses=(444,), max_tasks=100)
    catalogs = engine.load_catalogs(path, config)
    assert catalogs.swa_potential_gain[(100, 10)] == 3.2
