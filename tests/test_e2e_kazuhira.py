"""Prueba de punta a punta de execute_planning con Kazuhira activo.

Existe porque compilar y probar piezas sueltas no detecta errores de
tiempo de ejecución en la orquestación (variables fuera de alcance, un
archivo que no llega al zip, etc.)."""

from datetime import date

import modelo_abasto as engine
from tests._streamlit_stub import install as _install_streamlit_stub

_install_streamlit_stub()

import modules.les_enfants_terribles as m  # noqa: E402
from tests.e2e_fixture import (  # noqa: E402
    FakeUpload, build_plan_csv_bytes, build_workbook_bytes, read_zip_names,
)

STORES = {100: "CDMX", 101: "CDMX", 300: "CDMX"}   # 300 solo está en SCHEDULE
SKUS = [10, 11, 12]


def _run(**overrides):
    workbook = build_workbook_bytes(
        STORES, SKUS, schedule_only_stores={300},
        destination_stock={(101, 12): 50.0},        # tienda 101 / sku 12: sana
    )
    plan = build_plan_csv_bytes([(100, 10, 5), (101, 10, 5)])
    kwargs = dict(
        uploaded_copernico=[],
        uploaded_plans=[FakeUpload("plan.csv", plan)],
        database_bytes=workbook,
        origins=(444,),
        max_tasks=1000,
        run_date=date(2026, 10, 5),
        include_insumos=False,
        include_kazuhira_engine=True,
        enable_closed_stores_rule=False,
    )
    kwargs.update(overrides)
    return m.execute_planning(**kwargs)


def _bulk_lines(run):
    import csv, io, zipfile
    with zipfile.ZipFile(run["zip"]) as archive:
        text = archive.read("BulkCD_444.csv").decode("utf-8-sig")
    return {
        (int(r["WAREHOUSE_DESTINATION"]), int(r["RETAIL_ID"])): int(r["QUANTITY"])
        for r in csv.DictReader(io.StringIO(text))
    }


def test_kazuhira_end_to_end_covers_catalog_gaps_and_schedule_only_store():
    run = _run()
    k = run["kazuhira"]
    assert k["universe_stores"] == 3
    assert k["universe_from_fountain9"] == 2
    assert k["universe_schedule_only"] == 1
    lines = _bulk_lines(run)
    # La tienda 300 NO viene en Fountain9, solo en SCHEDULE: recibe sus 3 SKUs.
    # Duration/Lead Time heredados de la ciudad (5 + 3) con ADU 1 => 8.
    assert {sku: lines[(300, sku)] for sku in SKUS} == {10: 8, 11: 8, 12: 8}
    # SKUs que Fountain9 no pidió en tiendas que sí planea:
    assert lines[(100, 11)] == 8 and lines[(100, 12)] == 8 and lines[(101, 11)] == 8
    # Lo que Fountain9 sí pidió se respeta (Kazuhira no lo duplica):
    assert lines[(100, 10)] == 5 and lines[(101, 10)] == 5
    # (101, 12) tiene stock 50: sano, no se manda nada.
    assert (101, 12) not in lines


def test_healthy_universe_goes_to_csv_in_zip_not_to_base_rows():
    run = _run()
    names = read_zip_names(run["zip"])
    assert any("Universo_Catalogo_Sin_Necesidad" in n for n in names)
    healthy = run["status_counts"].get(m.CATALOG_UNIVERSE_HEALTHY_CUT, 0)
    assert healthy >= 1                       # (101, 12) al menos


def test_every_planned_store_sku_is_covered_or_declared_or_healthy():
    """La garantía: nada de las tiendas planeadas queda sin explicación."""
    run = _run()
    counts = run["status_counts"]
    total_combinations = len(STORES) * len(SKUS)      # 9
    accounted = sum(
        counts.get(label, 0)
        for label in (
            m.KAZUHIRA_CUT, m.CATALOG_UNIVERSE_HEALTHY_CUT,
            *m.KAZUHIRA_UNCOVERED_LABELS.values(),
            m.CATALOG_UNIVERSE_POST_KAZUHIRA_CUT,
            "OK COMPLETO POR FOUNTAIN9", "SIN RECOMENDACIÓN",
            "OK MANUAL POR FORECAST Y STOCK EN CERO",
        )
    )
    assert accounted == total_combinations


def test_run_without_kazuhira_has_no_healthy_csv():
    run = _run(include_kazuhira_engine=False, include_avl_fill=False)
    names = read_zip_names(run["zip"])
    assert not any("Universo_Catalogo_Sin_Necesidad" in n for n in names)


# --- explain_store_sku: "¿por qué no salió X en la tienda Y?" ------------------

def test_explain_sent_declared_healthy_and_no_trace():
    run = _run(max_tasks=4)          # presupuesto corto => habrá huecos declarados
    sent = m.explain_store_sku(run["zip"], 300, 10)
    healthy = m.explain_store_sku(run["zip"], 101, 12)
    nothing = m.explain_store_sku(run["zip"], 999, 10)
    assert healthy["verdict"].startswith("SANO") and healthy["healthy"]["STOCK"] == "50.0"
    assert nothing["verdict"].startswith("SIN RASTRO")
    assert "Kazuhira" in nothing["verdict"]
    # Con 4 tareas: 2 de Fountain9 + 2 de Kazuhira; el resto debe quedar DECLARADO.
    all_keys = [(s_, k) for s_ in STORES for k in SKUS]
    verdicts = {key: m.explain_store_sku(run["zip"], *key) for key in all_keys}
    assert not any(v["verdict"].startswith("SIN RASTRO") for v in verdicts.values())
    declared = [v for v in verdicts.values() if not v["sent"] and v["declared"]]
    assert declared, "con presupuesto corto debe haber huecos declarados"
    labels = {d["TIPO_DE_CORTE"] for v in declared for d in v["declared"]}
    assert m.KAZUHIRA_UNCOVERED_LABELS["SIN_TAREAS"] in labels


def test_explain_reports_units_and_reason_for_a_sent_line():
    run = _run()
    info = m.explain_store_sku(run["zip"], 300, 11)
    assert info["verdict"].startswith("SE ENVIÓ: 8")
    assert info["sent"][0]["MOTIVO"] == m.PLANNING_REASON_KAZUHIRA


# --- incoming: la columna real de DATA_TRANSFERS es INCOMING_TR --------------------

def _load_catalogs_from(workbook_bytes, tmp_path):
    path = tmp_path / "dt.xlsx"
    path.write_bytes(workbook_bytes)
    config = engine.Config(origin_warehouses=(444,), max_tasks=100)
    return engine.load_catalogs(path, config)


def test_incoming_tr_column_is_read(tmp_path):
    wb = build_workbook_bytes(STORES, SKUS, incoming={(100, 11): 7.0}, incoming_column="INCOMING_TR")
    catalogs = _load_catalogs_from(wb, tmp_path)
    assert catalogs.incoming_stock[(100, 11)] == 7.0


def test_legacy_incoming_column_name_still_works(tmp_path):
    wb = build_workbook_bytes(STORES, SKUS, incoming={(100, 11): 4.0}, incoming_column="INCOMING")
    catalogs = _load_catalogs_from(wb, tmp_path)
    assert catalogs.incoming_stock[(100, 11)] == 4.0


def test_missing_incoming_column_means_no_incoming(tmp_path):
    wb = build_workbook_bytes(STORES, SKUS, incoming_column=None)
    catalogs = _load_catalogs_from(wb, tmp_path)
    assert len(catalogs.incoming_stock) == 0


def test_incoming_suppresses_kazuhira_end_to_end():
    """(100, 11) tiene 3 de incoming con ADU 1 => 3 DOH: no hace falta
    enviar. Antes de leer INCOMING_TR se enviaban 8 unidades de más."""
    run = _run(
        database_bytes=build_workbook_bytes(
            STORES, SKUS, schedule_only_stores={300},
            destination_stock={(101, 12): 50.0}, incoming={(100, 11): 3.0},
        )
    )
    lines = _bulk_lines(run)
    assert (100, 11) not in lines
    assert lines[(100, 12)] == 8        # sin incoming sigue cubriéndose


# --- sin fila en STOCK => stock 0 e incoming 0, en todos los engines -----------------

def test_avl_alone_covers_catalog_combination_missing_from_stock_sheet():
    """Sin Kazuhira: antes AVL saltaba una combinación de CATALOGO sin fila
    en STOCK. Ahora la toma con stock 0 y la cubre."""
    workbook = build_workbook_bytes(
        STORES, SKUS, omit_stock_rows={(100, 12)},
    )
    run = _run(
        database_bytes=workbook, include_kazuhira_engine=False,
        include_avl_fill=True, avl_doh=3.0,
    )
    lines = _bulk_lines(run)
    assert (100, 12) in lines and lines[(100, 12)] > 0
    assert any(
        "sin fila en STOCK" in w and "AVL" in w for w in run["warnings"]
    ), run["warnings"][-6:]


def test_missing_stock_row_is_not_reported_as_skipped_anywhere():
    workbook = build_workbook_bytes(STORES, SKUS, omit_stock_rows={(100, 12)})
    run = _run(database_bytes=workbook)
    assert run["kazuhira"]["skipped_missing_stock"] == 0
    assert run["kazuhira"]["missing_stock_treated_as_zero"] == 1


def test_run_is_stamped_with_the_engine_build():
    run = _run()
    assert run["build"] == m.APP_BUILD
    assert run["warnings"][0] == f"Versión del motor: {m.APP_BUILD}."


def test_excel_carries_motivo_column_and_build_stamp():
    """Dos huellas para identificar con qué código se generó un Excel: la
    columna MOTIVO_KAZUHIRA en BASE_TRANSFERS (desde la v4, siempre que
    Kazuhira corre) y el sello de versión en RESUMEN."""
    import io, zipfile
    import openpyxl
    run = _run()
    with zipfile.ZipFile(run["zip"]) as archive:
        name = next(n for n in archive.namelist() if n.endswith(".xlsx"))
        workbook = openpyxl.load_workbook(
            io.BytesIO(archive.read(name)), read_only=True, data_only=True
        )
        header = next(workbook["BASE_TRANSFERS"].iter_rows(values_only=True))
        resumen = " | ".join(
            str(cell) for row in workbook["RESUMEN"].iter_rows(values_only=True)
            for cell in row if cell
        )
    assert "MOTIVO_KAZUHIRA" in header
    assert m.APP_BUILD in resumen
