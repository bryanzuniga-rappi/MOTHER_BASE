"""Pruebas del mecanismo de cobertura de quiebres sin fila de Fountain9 (apply_avl_fill
con candidate_mode="no_fountain9_coverage").
"""

import csv
from pathlib import Path
from types import SimpleNamespace

import modelo_abasto as engine
from tests._streamlit_stub import install as _install_streamlit_stub

_install_streamlit_stub()

import modules.les_enfants_terribles as m  # noqa: E402


def make_catalogs(**overrides) -> engine.Catalogs:
    base = dict(
        volume_m3={10: 1.0},
        blocked_products=set(),
        route_cost_blocks=set(),
        store_priority={100: 1},
        high_value={},
        rackeados_444=set(),
        store_capacity={100: 100.0},
        copernico_unusable_444={},
        unavailable_stock={},
        stock_base={(444, 10): 100.0, (100, 10): 0.0},
        golden_infaltables=set(),
        stores={
            444: {"city": "CDMX", "city_norm": "CDMX", "warehouse_name": "O444"},
            100: {"city": "CDMX", "city_norm": "CDMX", "warehouse_name": "STORE"},
        },
        storage={},
        warnings=[],
    )
    base.update(overrides)
    return engine.Catalogs(**base)


def make_result(**overrides) -> SimpleNamespace:
    base = dict(
        base_rows=[], allocation_rows=[], capacity_rows=[], tasks_used=0,
        max_tasks=100, warnings=[],
    )
    base.update(overrides)
    return SimpleNamespace(**base)


def catalog_rows(adu_10=1.0):
    return [{"WAREHOUSE_DESTINATION": 100, "RETAIL_ID": 10, "ADU": adu_10}]


CONFIG = engine.Config(origin_warehouses=(444,), max_tasks=100)


# --- apply_avl_fill(candidate_mode="no_fountain9_coverage") --------------

def test_basic_coverage_when_sku_absent_from_fountain9():
    """SKU quebrado (stock=0), sin fila de Fountain9, con stock en origen:
    debe cubrirse con ADU x (duration+leadtime)."""
    catalogs = make_catalogs()
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),  # SKU 10 no está en Fountain9
        duration_mode_by_store={100: 5.0},
        lead_time_mode_by_store={100: 3.0},
    )
    # ADU=2.0 x (5+3) = 16
    assert summary["units_added"] == 16
    assert summary["cases_sent"] == 1
    rows = [r for r in result.allocation_rows if r["RETAIL_ID"] == 10]
    assert sum(r["QUANTITY"] for r in rows) == 16


def test_skips_key_explicitly_in_excluded_keys():
    """Mecánica genérica: cualquier llave en excluded_keys se salta. Desde la sesión de
    "sin recomendación como quiebre real", quien llama a esta función decide
    exactamente qué va en excluded_keys.
    """
    catalogs = make_catalogs()
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys={(100, 10)},
        duration_mode_by_store={100: 5.0},
        lead_time_mode_by_store={100: 3.0},
    )
    assert summary["units_added"] == 0
    assert summary["cases_sent"] == 0


def test_sin_recomendacion_real_stockout_now_eligible():
    """Un "sin recomendación" de Fountain9 (CANTIDAD_OBJETIVO=0, por lo tanto NO entra
    en fountain_recommended_keys) cuyo stock_base real es 0 debe ser candidato
    elegible aquí.
    """
    catalogs = make_catalogs()  # stock_base[(100,10)] = 0 por el fixture
    result = make_result()
    # Excluded_keys = fountain_recommended_keys simulado: vacío, porque esta
    # tienda-SKU tuvo CANTIDAD_OBJETIVO=0 (sin recomendación), no entra en ese
    # set aunque SÍ tuvo fila en Fountain9.
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={100: 5.0},
        lead_time_mode_by_store={100: 3.0},
    )
    assert summary["units_added"] > 0
    assert summary["cases_sent"] == 1


def test_skips_when_destination_stock_positive():
    """No está quebrado (stock > 0): no debe cubrirse por este mecanismo."""
    catalogs = make_catalogs(stock_base={(444, 10): 100.0, (100, 10): 5.0})
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={100: 5.0},
        lead_time_mode_by_store={100: 3.0},
    )
    assert summary["units_added"] == 0
    assert summary.get("no_fountain9_candidates", 0) == 0


def test_skips_when_store_has_no_duration_leadtime_mode():
    """La tienda 100 no tiene moda de Duration/Lead Time (no está en los
    diccionarios) -> no hay base para calcular, se salta sin tronar."""
    catalogs = make_catalogs()
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={},  # vacío a propósito
        lead_time_mode_by_store={},
    )
    assert summary["units_added"] == 0
    assert summary.get("skipped_no_duration_data", 0) == 1


def test_fictitious_adu_used_when_adu_is_zero():
    """Sin ADU (ni propio ni de ciudad): usa el ADU ficticio de 0.14/día."""
    catalogs = make_catalogs()
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=0.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={100: 10.0},
        lead_time_mode_by_store={100: 4.0},
    )
    # 0.14 * (10+4) = 1.96 -> ceil = 2, pero el mínimo (default Config=3) gana
    assert summary["units_added"] == 3


def test_incoming_reduces_target():
    """STOCK.INCOMING debe restarse del objetivo bruto antes del mínimo."""
    catalogs = make_catalogs(incoming_stock={(100, 10): 10.0})
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={100: 5.0},
        lead_time_mode_by_store={100: 3.0},
    )
    # ADU=2 x (5+3) = 16, menos 10 de incoming = 6
    assert summary["units_added"] == 6


def test_incoming_fully_covers_need_skips_shipment():
    """Si el incoming ya cubre todo el objetivo, no se manda nada (no hay
    que sobre-enviar)."""
    catalogs = make_catalogs(incoming_stock={(100, 10): 999.0})
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={100: 5.0},
        lead_time_mode_by_store={100: 3.0},
    )
    assert summary["units_added"] == 0
    assert summary.get("skipped_covered_by_incoming", 0) == 1


def test_minimum_quantity_floor_applied():
    """Objetivo bruto chico: se eleva al mínimo configurado."""
    catalogs = make_catalogs()
    config = engine.Config(
        origin_warehouses=(444,), max_tasks=100, minimum_positive_quantity=5
    )
    result = make_result()
    summary = m.apply_avl_fill(
        result, catalog_rows(adu_10=0.1), catalogs, config, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={100: 2.0},
        lead_time_mode_by_store={100: 1.0},
    )
    # 0.1 * 3 = 0.3 -> ceil = 1, pero el mínimo configurado es 5
    assert summary["units_added"] == 5


def test_no_fountain9_cut_registered_in_breakdown_order():
    """Regresión: NO_FOUNTAIN9_CUT debe aparecer en BREAKDOWN_ORDER (si no, el overview
    no lo muestra con orden correcto).
    """
    assert m.NO_FOUNTAIN9_CUT in m.BREAKDOWN_ORDER


def test_no_fountain9_cut_attributed_to_its_own_engine():
    """Sus casos se atribuyen a Solidus, con la cobertura propia."""
    assert m.attribute_engine(m.NO_FOUNTAIN9_CUT) == "Solidus"
    assert m.attribute_row({"TIPO_DE_CORTE": m.NO_FOUNTAIN9_CUT}) == (
        "Solidus", "Cobertura sin Fountain9",
    )


def test_regla_demanda_and_tipo_de_corte_labels():
    catalogs = make_catalogs()
    result = make_result()
    m.apply_avl_fill(
        result, catalog_rows(adu_10=2.0), catalogs, CONFIG, set(), (), 1.0,
        candidate_mode="no_fountain9_coverage",
        excluded_keys=set(),
        duration_mode_by_store={100: 5.0},
        lead_time_mode_by_store={100: 3.0},
    )
    base_row = result.base_rows[0]
    assert base_row["REGLA_DEMANDA"] == "COBERTURA_SIN_FOUNTAIN9"
    assert base_row["TIPO_DE_CORTE"] == m.NO_FOUNTAIN9_CUT
    allocation = result.allocation_rows[0]
    assert allocation[m.PLANNING_REASON_COLUMN] == m.PLANNING_REASON_NO_FOUNTAIN9


# --- Duration/Lead Time moda en consolidate_plan_files --------------------

def _write_plan_csv(path: Path, headers: list[str], rows: list[list]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)


def test_duration_lead_time_mode_computed_per_store(tmp_path):
    plan_path = tmp_path / "plan.csv"
    _write_plan_csv(
        plan_path,
        [
            "Warehouseid", "SKU ID", "Predicted Demand for selected duration",
            "Predicted Opening Inventory",
            "Replenishment Quantity for Plan Duration (MOV)",
            "Net Inter-Store Transfers", "Duration",
            "Primary Source Lead Time (Days)",
        ],
        [
            [100, 10, "0", "0", "5", "0", "14", "3"],
            [100, 11, "0", "0", "5", "0", "14", "3"],
            [100, 12, "0", "0", "5", "0", "21", "5"],  # minoría
            [200, 20, "0", "0", "5", "0", "7", "2"],
        ],
    )
    catalogs = engine.Catalogs(
        volume_m3={}, blocked_products=set(), route_cost_blocks=set(),
        store_priority={}, high_value={}, rackeados_444=set(),
        store_capacity={}, copernico_unusable_444={}, unavailable_stock={},
        stock_base={}, golden_infaltables=set(), stores={}, storage={},
        warnings=[],
    )
    output_path = tmp_path / "canonical.csv"
    _, _consolidated, summary = m.consolidate_plan_files(
        [plan_path], catalogs, CONFIG, output_path
    )
    assert summary["duration_mode_by_store"][100] == 14.0  # moda: 14 gana 2 a 1
    assert summary["lead_time_mode_by_store"][100] == 3.0
    assert summary["duration_mode_by_store"][200] == 7.0
    assert summary["lead_time_mode_by_store"][200] == 2.0


def test_duration_lead_time_absent_when_columns_missing(tmp_path):
    """Si el archivo no trae esas columnas, los diccionarios de moda deben
    quedar vacíos, sin tronar la consolidación."""
    plan_path = tmp_path / "plan.csv"
    _write_plan_csv(
        plan_path,
        [
            "Warehouseid", "SKU ID", "Predicted Demand for selected duration",
            "Predicted Opening Inventory",
            "Replenishment Quantity for Plan Duration (MOV)",
            "Net Inter-Store Transfers",
        ],
        [[100, 10, "0", "0", "5", "0"]],
    )
    catalogs = engine.Catalogs(
        volume_m3={}, blocked_products=set(), route_cost_blocks=set(),
        store_priority={}, high_value={}, rackeados_444=set(),
        store_capacity={}, copernico_unusable_444={}, unavailable_stock={},
        stock_base={}, golden_infaltables=set(), stores={}, storage={},
        warnings=[],
    )
    output_path = tmp_path / "canonical.csv"
    _, _consolidated, summary = m.consolidate_plan_files(
        [plan_path], catalogs, CONFIG, output_path
    )
    assert summary["duration_mode_by_store"] == {}
    assert summary["lead_time_mode_by_store"] == {}


# --- STOCK.INCOMING (activación de la columna reservada) -----------------

def test_stock_incoming_loaded_when_present(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("STOCK")
    ws.append(["WAREHOUSE_ID", "PRODUCT_ID", "STOCK_DISPONIBLE_FINAL", "INCOMING"])
    ws.append([100, 10, 0, 7])

    # Hojas mínimas restantes requeridas por load_catalogs.
    for name, headers in [
        ("VOLUMETRIA", ["SKU", "PALLETS"]),
        ("BLOQUEOS_FORANEAS", ["SKU"]),
        ("BLOQUEOS", ["PRODUCT_ID"]),
        ("RUTA_COSTOS", ["Destination", "Catalog ID"]),
        ("PRIORIDAD", ["WAREHOUSE_ID", "PRIORIDAD"]),
        ("444_HV", ["EAN", "Category"]),
        ("831_HV", ["EAN", "Category"]),
        ("RACKEADOS", ["WHS", "SYNC"]),
        ("CAP_RECIBO", ["WH_ID", "CAP"]),
        ("CATALOGO", ["WAREHOUSE_ID", "PRODUCT_ID", "ADU"]),
        ("KVI", ["WAREHOUSE_ID", "PRODUCT_ID", "KVI"]),
        ("SHARE_VENTAS", ["WAREHOUSE_ID", "SHARE"]),
        ("NO_DISPONIBLE", ["WAREHOUSE_ID", "PRODUCT_ID", "STOCK"]),
        ("POR_MERMAR", ["WAREHOUSE_ID", "PRODUCT_ID", "STOCK_AVAILABLE", "VALUE_STOCK", "ARRIVAL_DATE", "EXPIRATION_DATE"]),
        ("OWNER", ["WAREHOUSE_ID", "PRODUCT_ID", "OWNER_NAME", "STOCK_DISPONIBLE_FINAL"]),
        ("INSUMOS", ["WAREHOUSE_DESTINATION", "WAREHOUSE_SOURCE", "RETAIL_ID", "QUANTITY", "PLANNED_DATE", "ROUTE", "DELIVERY_PRIORITY"]),
        ("GOLDEN_INFALTABLES_ANCHOR", ["WAREHOUSE_ID", "PRODUCT_ID_SYNC", "IS_INFALTABLE", "IS_GOLDEN", "IS_ANCHOR"]),
        ("TIENDA", ["CITY", "WAREHOUSE_ID", "WAREHOUSE_NAME"]),
        ("STORAGE", ["PRODUCT_ID", "STORAGE_NAME"]),
        ("TIENDAS_CERRADAS", ["WAREHOUSE_ID"]),
        ("SCHEDULE", ["CITY", "WAREHOUSE_ID", "WAREHOUSE_NAME", "ORIGEN", "DAYS"]),
    ]:
        extra = wb.create_sheet(name)
        extra.append(headers)
        if name == "TIENDA":
            extra.append(["Ciudad de México", 100, "STORE"])
            extra.append(["Ciudad de México", 444, "O444"])

    xlsx_path = tmp_path / "DATA_TRANSFERS.xlsx"
    wb.save(xlsx_path)

    catalogs = engine.load_catalogs(xlsx_path, CONFIG)
    assert catalogs.incoming_stock[(100, 10)] == 7.0


def test_stock_incoming_empty_when_column_absent(tmp_path):
    """Sin la columna INCOMING (el caso normal hoy): incoming_stock debe
    quedar vacío, sin romper la carga de STOCK."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("STOCK")
    ws.append(["WAREHOUSE_ID", "PRODUCT_ID", "STOCK_DISPONIBLE_FINAL"])
    ws.append([100, 10, 5])
    for name, headers in [
        ("VOLUMETRIA", ["SKU", "PALLETS"]),
        ("BLOQUEOS_FORANEAS", ["SKU"]),
        ("BLOQUEOS", ["PRODUCT_ID"]),
        ("RUTA_COSTOS", ["Destination", "Catalog ID"]),
        ("PRIORIDAD", ["WAREHOUSE_ID", "PRIORIDAD"]),
        ("444_HV", ["EAN", "Category"]),
        ("831_HV", ["EAN", "Category"]),
        ("RACKEADOS", ["WHS", "SYNC"]),
        ("CAP_RECIBO", ["WH_ID", "CAP"]),
        ("CATALOGO", ["WAREHOUSE_ID", "PRODUCT_ID", "ADU"]),
        ("KVI", ["WAREHOUSE_ID", "PRODUCT_ID", "KVI"]),
        ("SHARE_VENTAS", ["WAREHOUSE_ID", "SHARE"]),
        ("NO_DISPONIBLE", ["WAREHOUSE_ID", "PRODUCT_ID", "STOCK"]),
        ("POR_MERMAR", ["WAREHOUSE_ID", "PRODUCT_ID", "STOCK_AVAILABLE", "VALUE_STOCK", "ARRIVAL_DATE", "EXPIRATION_DATE"]),
        ("OWNER", ["WAREHOUSE_ID", "PRODUCT_ID", "OWNER_NAME", "STOCK_DISPONIBLE_FINAL"]),
        ("INSUMOS", ["WAREHOUSE_DESTINATION", "WAREHOUSE_SOURCE", "RETAIL_ID", "QUANTITY", "PLANNED_DATE", "ROUTE", "DELIVERY_PRIORITY"]),
        ("GOLDEN_INFALTABLES_ANCHOR", ["WAREHOUSE_ID", "PRODUCT_ID_SYNC", "IS_INFALTABLE", "IS_GOLDEN", "IS_ANCHOR"]),
        ("TIENDA", ["CITY", "WAREHOUSE_ID", "WAREHOUSE_NAME"]),
        ("STORAGE", ["PRODUCT_ID", "STORAGE_NAME"]),
        ("TIENDAS_CERRADAS", ["WAREHOUSE_ID"]),
        ("SCHEDULE", ["CITY", "WAREHOUSE_ID", "WAREHOUSE_NAME", "ORIGEN", "DAYS"]),
    ]:
        extra = wb.create_sheet(name)
        extra.append(headers)
        if name == "TIENDA":
            extra.append(["Ciudad de México", 100, "STORE"])
            extra.append(["Ciudad de México", 444, "O444"])
    xlsx_path = tmp_path / "DATA_TRANSFERS.xlsx"
    wb.save(xlsx_path)

    catalogs = engine.load_catalogs(xlsx_path, CONFIG)
    assert catalogs.incoming_stock == {}
    assert catalogs.stock_base[(100, 10)] == 5.0
